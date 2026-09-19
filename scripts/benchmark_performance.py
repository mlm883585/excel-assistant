"""Synthetic read fast-path benchmarks: DuckDB CSV vs pandas, single vs double xlsx pass.

Run: .venv\\Scripts\\python.exe scripts\\benchmark_performance.py [--rows N]
Writes build/performance-metrics.json. Does not touch any source file.
"""
import argparse
import json
from pathlib import Path
import sys
import time
import uuid
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


def measure(timings, name, fn):
    start = time.perf_counter()
    value = fn()
    timings[name] = round(time.perf_counter() - start, 3)
    print(name, timings[name], flush=True)
    return value


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--rows", type=int, default=100000)
    args = parser.parse_args()

    from openpyxl import Workbook
    from assistant.store import Store
    from assistant.tables import read, _read_csv_fast, _xlsx_has_formula
    from assistant.models import InputSelection

    folder = ROOT / "build" / ("performance-benchmark-" + uuid.uuid4().hex)
    folder.mkdir(parents=True)
    store = Store(folder / "state")
    task = store.create()["id"]
    timings = {}

    # CSV: rows x 3 columns, leading-zero code plus a legal NA text.
    csv_path = folder / "large.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
        stream.write("code,qty,note\n")
        for i in range(args.rows):
            stream.write(f"{i:08d},{i},NA\n")
    with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 0):
        fast = measure(timings, "csv_read_duckdb_seconds", lambda: _read_csv_fast(csv_path))
    slow = measure(
        timings,
        "csv_read_pandas_seconds",
        lambda: __import__("pandas").read_csv(csv_path, header=None, dtype=str, keep_default_na=False, encoding="utf-8-sig", skip_blank_lines=False),
    )
    assert fast is not None and len(fast) == len(slow) == args.rows + 1, (len(fast), len(slow))
    assert fast.iloc[1, 0] == slow.iloc[1, 0] == "00000000"

    # Value-only xlsx (no formulas): single-pass vs forced double-pass.
    xlsx_path = folder / "plain.xlsx"
    book = Workbook()
    book.active.append(["code", "qty", "note"])
    for i in range(args.rows):
        book.active.append([f"{i:08d}", i, "NA"])
    book.save(xlsx_path)
    book.close()
    assert not _xlsx_has_formula(xlsx_path)
    file_id = store.import_file(task, xlsx_path)["id"]
    selection = InputSelection(file_id=file_id, header_row=1, sheet=0)
    single = measure(timings, "xlsx_read_single_pass_seconds", lambda: read(store, task, selection))
    with mock.patch("assistant.tables._xlsx_has_formula", return_value=True):
        double = measure(timings, "xlsx_read_double_pass_seconds", lambda: read(store, task, selection))
    assert len(single) == len(double) == args.rows
    assert single.code.iloc[0] == double.code.iloc[0] == "00000000"
    assert single["__source_row"].tolist() == double["__source_row"].tolist()

    results = {"timings": timings, "rows": args.rows, "folder": str(folder)}
    (ROOT / "build/performance-metrics.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(results, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
