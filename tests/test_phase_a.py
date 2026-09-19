import asyncio
import json
import tempfile
from pathlib import Path
import unittest

import pandas as pd

from data_toolkit.models import ToolkitError
from data_toolkit.profiling import profile_columns
from data_toolkit.sql_runner import run_query
from assistant.store import Store
from assistant.models import Operation, InputSelection
from assistant.tables import run_operation, read, make_report
from assistant.workbook_model import blank_workbook
from assistant.workbook_operations import edit_workbook
from assistant.workbooks import Workbooks


class PhaseATests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / 'state')
        self.task = self.store.create()['id']
        self.books = Workbooks(self.store)

    def imported(self, name, text):
        path = self.root / name
        path.write_text(text, encoding='utf-8-sig')
        return self.store.import_file(self.task, path)['id']

    def test_profile_columns_json_safe_and_stats(self):
        frame = pd.DataFrame({'物料': ['螺丝', '螺母', '螺丝', '垫片'], '数量': [10, 20, 30, None], '单价': [1.5, 2.0, 1.5, 3.0]})
        cols = profile_columns(frame, 3)
        json.dumps(cols)  # must not raise (no np scalars leaking)
        by_name = {c['name']: c for c in cols}
        self.assertEqual(by_name['数量']['null_rate'], 0.25)
        self.assertEqual(by_name['物料']['unique'], 3)
        self.assertEqual(by_name['单价']['samples'], [1.5, 2.0, 3.0])
        for c in cols:
            self.assertIsInstance(c['name'], str)
            self.assertIsInstance(c['dtype'], str)
            self.assertIsInstance(c['null_rate'], float)
            self.assertIsInstance(c['unique'], int)
            for v in c['samples'] + [c['min'], c['max']]:
                self.assertIsInstance(v, (str, int, float, bool, type(None)))

    def test_set_formula_writes_pending_and_clears(self):
        book = blank_workbook()
        sid = book['sheets'][0]['id']
        result = edit_workbook(book, sid, {'edits': [{'kind': 'set_formula', 'range': {'r0': 0, 'c0': 2, 'r1': 0, 'c1': 2}, 'formulas': [['=A1*2']]}]})
        cell = result['sheets'][0]['cells']['0,2']
        self.assertEqual(cell['formula'], '=A1*2')
        self.assertEqual(cell['result_state'], 'pending')
        self.assertIsNone(cell['value'])
        book['sheets'][0]['cells']['0,0'] = {'value': 5, 'type': 'number'}
        cleared = edit_workbook(book, sid, {'edits': [{'kind': 'set_formula', 'range': {'r0': 0, 'c0': 0, 'r1': 0, 'c1': 0}, 'formulas': [['']]}]})
        self.assertNotIn('0,0', cleared['sheets'][0]['cells'])

    def test_set_formula_rejects_invalid(self):
        book = blank_workbook()
        sid = book['sheets'][0]['id']
        with self.assertRaisesRegex(ValueError, '暂不支持公式函数'):
            edit_workbook(book, sid, {'edits': [{'kind': 'set_formula', 'range': {'r0': 0, 'c0': 0, 'r1': 0, 'c1': 0}, 'formulas': [['=XLOOKUP(A1,B:B,C:C)']]}]})
        with self.assertRaisesRegex(ValueError, '数组'):
            edit_workbook(book, sid, {'edits': [{'kind': 'set_formula', 'range': {'r0': 0, 'c0': 0, 'r1': 0, 'c1': 0}, 'formulas': [['={1,2,3}']]}]})
        with self.assertRaisesRegex(ValueError, '尺寸'):
            edit_workbook(book, sid, {'edits': [{'kind': 'set_formula', 'range': {'r0': 0, 'c0': 0, 'r1': 1, 'c1': 1}, 'formulas': [['=A1']]}]})

    def test_set_formula_scope_enforced(self):
        record = self.books.open(self.task)
        book = record['snapshot']
        sid = book['sheets'][0]['id']
        op = {'kind': 'edit_workbook', 'inputs': [{'workbook_id': book['id'], 'version': 1, 'sheet_id': sid, 'range': {'r0': 0, 'c0': 0, 'r1': 0, 'c1': 0}}], 'params': {'edits': [{'kind': 'set_formula', 'range': {'r0': 1, 'c0': 0, 'r1': 1, 'c1': 0}, 'formulas': [['=A1*2']]}]}}
        with self.assertRaisesRegex(ValueError, '超出用户选区'):
            run_operation(self.store, self.task, op)

    def test_sql_readonly_enforced(self):
        frame = pd.DataFrame({'a': [1, 2], 'b': [3, 4]})
        self.assertEqual(run_query({'t': frame}, 'SELECT a FROM t WHERE b > 3')['a'].tolist(), [2])
        for bad in ('DELETE FROM t', 'UPDATE t SET a = 1', "SELECT * FROM read_csv_auto('x.csv')"):
            with self.assertRaises(ToolkitError):
                run_query({'t': frame}, bad)

    def test_coerce_numeric_columns_protects_leading_zero_codes(self):
        from data_toolkit.sql_runner import coerce_numeric_columns
        frame = pd.DataFrame({'编号': ['001', '002'], '数量': ['10', '20'], '金额': ['1.5', '']})
        result = coerce_numeric_columns(frame)
        self.assertEqual(result['编号'].tolist(), ['001', '002'])  # leading zeros preserved
        self.assertTrue(pd.api.types.is_numeric_dtype(result['数量']))
        self.assertEqual(result['金额'].iloc[0], 1.5)
        self.assertTrue(pd.isna(result['金额'].iloc[1]))

    def test_create_server_registers_all_tools(self):
        from assistant.mcp_server import create_server
        server = create_server(self.root / 'state', self.task)
        names = {t.name for t in asyncio.run(server.list_tools())}
        self.assertEqual(len(names), 13)
        for tool in ('datacraft_profile', 'datacraft_sql', 'datacraft_report', 'datacraft_audit', 'datacraft_formula_generate', 'datacraft_formula_explain', 'datacraft_chart', 'datacraft_pivot_table'):
            self.assertIn(tool, names)

    def test_report_operation_without_model(self):
        fid = self.imported('数据.csv', '物料,数量\n螺丝,10\n螺母,20\n螺丝,30\n')
        info = run_operation(self.store, self.task, {'kind': 'report', 'inputs': [{'file_id': fid}], 'params': {}})
        self.assertEqual(info['name'], '数据质量报告.xlsx')
        self.assertEqual(info['statistics']['行数'], 3)
        self.assertEqual(info['statistics']['字段数'], 2)
        self.assertEqual(info['statistics']['重复行'], 0)
        # 质量报告是独立产物，不产生候选核对记录，可直接打开/导出
        review = Workbooks(self.store).review(self.task, info['id'])
        self.assertTrue(review.get('legacy'))

    def test_report_operation_requires_single_file(self):
        with self.assertRaisesRegex(ValueError, '需要输入'):
            Operation.model_validate({'kind': 'report', 'inputs': [], 'params': {}})
        with self.assertRaisesRegex(ValueError, '一个明确输入'):
            Operation.model_validate({'kind': 'report', 'inputs': [{'file_id': 'a'}, {'file_id': 'b'}], 'params': {}})

    def test_real_xlsx_profile_sql_report_smoke(self):
        from openpyxl import Workbook
        from data_toolkit.sql_runner import coerce_numeric_columns
        path = self.root / '销售.xlsx'
        wb = Workbook()
        orders = wb.active
        orders.title = '订单'
        orders.append(['编号', '物料', '数量', '单价'])
        for row in ([1, '螺丝', 10, 1.5], [2, '螺母', 20, None], [3, '螺丝', 30, 1.5], [1, '螺丝', 10, 1.5]):
            orders.append(row)
        for r in range(2, 6):
            orders.cell(r, 1).number_format = '000'
        catalog = wb.create_sheet('物料')
        catalog.append(['物料', '标准单价'])
        catalog.append(['螺丝', 1.5])
        catalog.append(['螺母', 2.0])
        wb.save(path)
        fid = self.store.import_file(self.task, path)['id']

        orders_frame = read(self.store, self.task, InputSelection(file_id=fid, sheet='订单'))
        # 前导零编号保留为文本；单价有一处空值
        self.assertEqual(orders_frame['编号'].tolist(), ['001', '002', '003', '001'])
        data_cols = [c for c in orders_frame.columns if not str(c).startswith('__source_')]
        by_name = {c['name']: c for c in profile_columns(orders_frame[data_cols], 3)}
        self.assertEqual(by_name['单价']['null_rate'], 0.25)
        self.assertEqual(by_name['物料']['unique'], 2)
        self.assertNotIn('', by_name['物料']['samples'])

        catalog_frame = read(self.store, self.task, InputSelection(file_id=fid, sheet='物料'))
        joined = run_query(
            {'orders': coerce_numeric_columns(orders_frame[data_cols]), 'catalog': coerce_numeric_columns(catalog_frame[[c for c in catalog_frame.columns if not str(c).startswith('__source_')]])},
            'SELECT orders.物料, SUM(orders.数量) AS total FROM orders JOIN catalog ON orders.物料 = catalog.物料 GROUP BY orders.物料 ORDER BY total DESC',
        )
        self.assertEqual(joined['物料'].tolist(), ['螺丝', '螺母'])
        self.assertEqual(joined['total'].tolist(), [50, 20])

        info = make_report(self.store, self.task, orders_frame)
        self.assertEqual(info['statistics']['行数'], 4)
        self.assertEqual(info['statistics']['字段数'], 4)
        self.assertEqual(info['statistics']['空值单元格'], 1)
        self.assertEqual(info['statistics']['重复行'], 1)
