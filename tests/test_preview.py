import json
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import patch

from assistant.preview import PreviewService, stream_rows
from assistant.store import Store


class PreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = Store(self.root / '应用 数据')
        self.task = self.store.create()['id']
        self.service = PreviewService(self.store)

    def tearDown(self):
        self.service.close()
        self.temp.cleanup()

    def csv(self, content='编码,数量\n001,2\nNA,3\n', name='业务 表.csv'):
        path = self.root / name
        path.write_text(content, encoding='utf-8-sig')
        return self.store.import_file(self.task, path)

    def start(self, info, header=1):
        return self.service.start(self.task, {'file_id': info['id'], 'sheet': 0, 'header_row': header})['preview_id']

    def finish(self, key):
        deadline = time.monotonic() + 30
        while time.monotonic() < deadline:
            result = self.service.page(key)
            if result['state'] != 'indexing':
                return result
            time.sleep(.025)
        self.fail('preview timed out')

    def test_csv_pages_keep_codes_and_source_without_reparse(self):
        info = self.csv('编码,数量\n' + ''.join(f'{i:05d},NA\n' for i in range(450)))
        key = self.start(info)
        result = self.finish(key)
        self.assertEqual(result['total'], 450)
        with patch('assistant.preview.stream_rows', side_effect=AssertionError('must not reparse')):
            for offset in [400, 0, 200, 100]:
                page = self.service.page(key, offset)
                self.assertEqual(page['rows'][0]['编码'], f'{offset:05d}')
                self.assertEqual(page['rows'][0]['数量'], 'NA')
                self.assertEqual(page['rows'][0]['__source_row'], offset + 2)
            self.assertEqual(self.start(info), key)

    def test_partial_metadata_initialization_keeps_indexing(self):
        from unittest.mock import Mock
        info=self.csv();source,_=self.store.resolve_file(self.task,info['id'])
        partial=self.root/'initial.partial';final=self.root/'initial.sqlite';error=self.root/'initial.error.json'
        with sqlite3.connect(partial) as db:
            db.execute('CREATE TABLE metadata(columns TEXT, count INTEGER, complete INTEGER)')
        db.close()
        key='initializing'
        self.service.requests[key]=(final,partial,error,source,[source.stat().st_size,source.stat().st_mtime_ns])
        self.service.active=key
        with patch.object(self.service,'process',Mock(is_alive=Mock(return_value=True))):
            self.assertEqual(self.service.page(key)['state'],'indexing')

    def test_xlsx_matches_business_reader_and_formula_failure(self):
        from openpyxl import Workbook
        from assistant.tables import preview
        path = self.root / 'test.xlsx'
        book = Workbook(); sheet = book.active
        sheet.append(['标题', '备注']); sheet.append(['编码', '数量']); sheet.append([7, 3]); sheet['A3'].number_format = '00000'; sheet.append(['NA', 8])
        book.save(path); book.close()
        info = self.store.import_file(self.task, path)
        result = self.finish(self.start(info, 2))
        expected = preview(self.store, self.task, {'file_id': info['id'], 'sheet': 0, 'header_row': 2})
        self.assertEqual(result['rows'], expected['rows'])
        self.assertEqual(result['columns'], expected['columns'])
        book = Workbook(); book.active.append(['编码']); book.active.append(['=1+1']); book.save(path); book.close()
        invalid = self.store.import_file(self.task, path)
        error = self.finish(self.start(invalid))
        self.assertEqual(error['state'], 'failed')
        self.assertIn('公式没有缓存值', error['message'])

    def test_cached_formula_and_dates_match(self):
        import zipfile
        from datetime import datetime
        import xlsxwriter
        from assistant.tables import preview
        path = self.root / 'cached.xlsx'
        with xlsxwriter.Workbook(path) as book:
            sheet = book.add_worksheet()
            sheet.write_row(0, 0, ['日期', '结果'])
            sheet.write_datetime(1, 0, datetime(2026, 1, 1), book.add_format({'num_format': 'yyyy-mm-dd'}))
            sheet.write_formula(1, 1, '=1+1', None, 2)
        info = self.store.import_file(self.task, path)
        result = self.finish(self.start(info))
        self.assertEqual(result['rows'][0]['结果'], 2)
        self.assertTrue(result['rows'][0]['日期'].startswith('2026-01-01T00:00:00'))

    def test_gb18030_quoted_newline_and_header(self):
        path = self.root / '旧编码.csv'
        path.write_text('标题,备注\n编码,描述\n001,"甲\n乙"\nNA,丙\n', encoding='gb18030', newline='')
        info = self.store.import_file(self.task, path)
        value = self.finish(self.start(info, 2))
        self.assertEqual(value['rows'][0]['描述'], '甲\n乙')
        self.assertEqual(value['rows'][1]['__source_row'], 4)

    def test_invalid_headers_and_reserved_columns(self):
        for content in ['a,a\n1,2\n', 'a,\n1,2\n', '__source_row,a\n1,2\n']:
            info = self.csv(content)
            self.assertEqual(self.finish(self.start(info))['state'], 'failed')

    def test_changed_file_and_corrupt_cache_rebuild(self):
        info = self.csv(); key = self.start(info); self.finish(key)
        final = self.service.requests[key][0]
        final.write_bytes(b'broken')
        self.assertEqual(self.service.page(key)['state'], 'failed')
        self.assertEqual(self.finish(self.start(info))['total'], 2)
        path, _ = self.store.resolve_file(self.task, info['id']); path.write_text('编码,数量\n009,1\n')
        self.assertEqual(self.service.page(key)['state'], 'stale')
        new = self.start(info); self.assertNotEqual(key, new)
        self.assertEqual(self.finish(new)['total'], 1)

    def test_duplicate_start_cancel_and_close(self):
        info = self.csv('编码,数量\n' + '001,3\n' * 100000)
        key = self.start(info); process = self.service.process
        self.assertEqual(key, self.start(info)); self.assertIs(process, self.service.process)
        self.assertFalse(self.service.cancel('another-id'))
        self.service.cancel(key)
        self.assertIsNone(self.service.process)
        self.assertFalse(list(self.store.root.glob('tasks/*/preview-cache/*.partial*')))
        self.start(info); self.service.close()
        self.assertIsNone(self.service.process)
        with self.assertRaises(ValueError): self.start(info)

    def test_quota_and_lru(self):
        self.service.quota = 100
        result = self.finish(self.start(self.csv()))
        self.assertEqual(result['state'], 'failed'); self.assertIn('容量', result['message'])
        self.service.quota = 50000
        first = self.start(self.csv()); self.finish(first); first_path = self.service.requests[first][0]
        second = self.start(self.csv(name='second.csv')); self.finish(second)
        self.assertFalse(first_path.exists())
        self.assertLessEqual(sum(p.stat().st_size for p in self.store.root.glob('tasks/*/preview-cache/*.sqlite')), self.service.quota)

    def test_summary_and_event_paging(self):
        with self.store.connect() as db:
            db.executemany('INSERT INTO tasks(id,body) VALUES (?,?)', [(f'{i:032x}', json.dumps({'status':'pending','files':[{'name':f'file{i}'}], 'outputs':[{'secret':'not-summary'}]})) for i in range(1000)])
            db.executemany('INSERT INTO events(task,body) VALUES (?,?)', [(self.task, json.dumps({'kind':'message','data':i})) for i in range(10000)])
        summary = self.store.summaries(); self.assertEqual(len(summary['items']), 30); self.assertTrue(summary['has_more'])
        self.assertNotIn('outputs', summary['items'][0])
        page = self.store.event_page(self.task); self.assertEqual(len(page['items']), 200)
        older = self.store.event_page(self.task, page['items'][0]['id']); self.assertLess(older['items'][-1]['id'], page['items'][0]['id'])
        with self.store.connect() as db:
            plan = db.execute('EXPLAIN QUERY PLAN SELECT id,body FROM events WHERE task=? AND id>? ORDER BY id LIMIT 200', (self.task, 0)).fetchall()
        self.assertIn('events_task_id', str(plan))

    def test_start_during_worker_exit_returns_retry(self):
        from api import api
        from unittest.mock import Mock
        jobs = Mock(); jobs.processes = {'task': Mock()}
        jobs.processes['task'].is_alive.return_value = True
        with patch.object(api, '_jobs', jobs), patch.object(api, 'previews') as service:
            self.assertEqual(api.preview_start(self.task, {})['state'], 'busy')
            service.assert_not_called()


if __name__ == '__main__':
    unittest.main()
