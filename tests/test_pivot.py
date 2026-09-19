import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook
from openpyxl.pivot.table import TableDefinition

from assistant.models import InputSelection, Operation
from assistant.pivot_excel import build_pivot_xlsx
from assistant.store import Store
from assistant.tables import STEP_LABELS, aggregate_pivot, emit_step, read, run_operation
from assistant.workbook_excel import import_workbook


class AggregatePivotTests(unittest.TestCase):
    def frame(self):
        return pd.DataFrame({
            '物料': ['螺丝', '螺母', '螺丝', '螺母', '螺丝'],
            '地区': ['东', '西', '东', '东', '西'],
            '数量': [10, 20, 30, 40, 5],
            '金额': [100, 200, 300, 400, 50],
        })

    def test_single_value_clean_columns(self):
        result = aggregate_pivot(self.frame(), ['物料'], '地区', ['数量'], 'sum')
        self.assertEqual(result.columns.tolist(), ['物料', '东', '西'])
        by_row = {r['物料']: r for r in result.to_dict('records')}
        self.assertEqual(by_row['螺丝']['东'], 40)
        self.assertEqual(by_row['螺母']['东'], 40)
        self.assertEqual(by_row['螺母']['西'], 20)

    def test_multi_value_flattened_columns(self):
        result = aggregate_pivot(self.frame(), ['物料'], '地区', ['数量', '金额'], 'sum')
        self.assertIn('数量 · 东', result.columns.tolist())
        self.assertIn('金额 · 西', result.columns.tolist())

    def test_empty_column(self):
        result = aggregate_pivot(self.frame(), ['物料'], None, ['数量'], 'sum')
        self.assertEqual(result.columns.tolist(), ['物料', '数量'])
        by_row = {r['物料']: r for r in result.to_dict('records')}
        self.assertEqual(by_row['螺丝']['数量'], 45)

    def test_multiple_rows(self):
        result = aggregate_pivot(self.frame(), ['物料', '地区'], None, ['数量'], 'sum')
        self.assertEqual(result.columns.tolist(), ['物料', '地区', '数量'])
        self.assertEqual(len(result), 4)

    def test_aggregate_variants(self):
        result = aggregate_pivot(self.frame(), ['物料'], None, ['数量'], 'count')
        by_row = {r['物料']: r for r in result.to_dict('records')}
        self.assertEqual(by_row['螺丝']['数量'], 3)

    def test_validation(self):
        with self.assertRaisesRegex(ValueError, '至少一个行字段'):
            aggregate_pivot(self.frame(), [], '地区', ['数量'], 'sum')
        with self.assertRaisesRegex(ValueError, '至少一个数值字段'):
            aggregate_pivot(self.frame(), ['物料'], '地区', [], 'sum')
        with self.assertRaisesRegex(ValueError, '存在的字段'):
            aggregate_pivot(self.frame(), ['物料'], '地区', ['不存在'], 'sum')
        with self.assertRaisesRegex(ValueError, '有效汇总规则'):
            aggregate_pivot(self.frame(), ['物料'], '地区', ['数量'], 'median')


class BuildPivotXlsxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / 'state')
        self.task = self.store.create()['id']

    def imported(self, name, text):
        path = self.root / name
        path.write_text(text, encoding='utf-8-sig')
        return self.store.import_file(self.task, path)['id']

    def frame(self, file_id):
        return read(self.store, self.task, InputSelection(file_id=file_id))

    def build(self, spec=None):
        fid = self.imported('数据.csv', '物料,地区,数量,金额\n螺丝,东,10,100\n螺母,西,20,200\n螺丝,东,30,300\n')
        target = self.root / 'out.xlsx'
        build_pivot_xlsx(self.frame(fid), spec or {'rows': ['物料'], 'columns': ['地区'], 'values': [{'field': '数量', 'aggregate': 'sum'}], 'name': '库存透视'}, target)
        return target

    def test_parts_and_content_types(self):
        target = self.build()
        with ZipFile(target) as archive:
            names = set(archive.namelist())
            for part in ('xl/pivotTables/pivotTable1.xml', 'xl/pivotCache/pivotCacheDefinition1.xml',
                         'xl/pivotCache/pivotCacheRecords1.xml', 'xl/worksheets/sheet2.xml'):
                self.assertIn(part, names, part)
            content_types = archive.read('[Content_Types].xml').decode('utf-8')
            self.assertIn('pivotTable+xml', content_types)
            self.assertIn('pivotCacheDefinition+xml', content_types)
            self.assertIn('pivotCacheRecords+xml', content_types)

    def test_openpyxl_reads_workbook(self):
        target = self.build()
        book = load_workbook(target, read_only=True)
        try:
            self.assertEqual(book.sheetnames, ['数据', '库存透视'])
        finally:
            book.close()

    def test_roundtrip_definition(self):
        target = self.build()
        with ZipFile(target) as archive:
            table = TableDefinition.from_tree(ET.fromstring(archive.read('xl/pivotTables/pivotTable1.xml')))
        self.assertEqual(table.name, '库存透视')
        self.assertEqual(table.cacheId, 1)
        self.assertEqual(table.dataCaption, '求和项')
        self.assertEqual(len(table.pivotFields), 3)
        self.assertEqual(table.pivotFields[0].axis, 'axisRow')
        self.assertEqual(table.pivotFields[1].axis, 'axisCol')
        self.assertIsNone(table.pivotFields[2].axis)
        self.assertTrue(table.pivotFields[2].dataField)
        self.assertEqual(table.rowFields[0].x, 0)
        self.assertEqual(table.colFields[0].x, 1)
        self.assertEqual(table.dataFields[0].fld, 2)
        self.assertEqual(table.dataFields[0].subtotal, 'sum')
        self.assertEqual(table.dataFields[0].name, '数量')
        self.assertTrue(table.location.ref)

    def test_readonly_on_reimport(self):
        target = self.build()
        _, issues = import_workbook(target)
        self.assertTrue(any('透视表' in str(i) for i in issues), f'预期透视表只读提示，实际: {issues}')

    def test_validation(self):
        fid = self.imported('数据.csv', '物料,数量\n螺丝,10\n螺母,20\n')
        frame = self.frame(fid)
        target = self.root / 'out.xlsx'
        with self.assertRaisesRegex(ValueError, '至少需要一个行字段'):
            build_pivot_xlsx(frame, {'rows': [], 'columns': [], 'values': [{'field': '数量', 'aggregate': 'sum'}]}, target)
        with self.assertRaisesRegex(ValueError, '最多支持一个列字段'):
            build_pivot_xlsx(frame, {'rows': ['物料'], 'columns': ['数量', '数量'], 'values': [{'field': '数量', 'aggregate': 'sum'}]}, target)
        with self.assertRaisesRegex(ValueError, '需要一个数值字段'):
            build_pivot_xlsx(frame, {'rows': ['物料'], 'columns': [], 'values': [], }, target)
        with self.assertRaisesRegex(ValueError, '有效汇总规则'):
            build_pivot_xlsx(frame, {'rows': ['物料'], 'columns': [], 'values': [{'field': '数量', 'aggregate': 'median'}]}, target)


class PivotTableOperationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / 'state')
        self.task = self.store.create()['id']

    def imported(self, name, text):
        path = self.root / name
        path.write_text(text, encoding='utf-8-sig')
        return self.store.import_file(self.task, path)['id']

    def test_dispatch_and_step_label(self):
        fid = self.imported('数据.csv', '物料,地区,数量\n螺丝,东,10\n螺母,西,20\n螺丝,东,30\n')
        info = run_operation(self.store, self.task, {'kind': 'pivot_table', 'inputs': [{'file_id': fid}], 'params': {'rows': ['物料'], 'columns': ['地区'], 'values': [{'field': '数量', 'aggregate': 'sum'}]}})
        self.assertEqual(info['name'], '透视表.xlsx')
        self.assertEqual(info['statistics']['input_rows'], 3)
        self.assertIn(info['id'], [o['id'] for o in self.store.get(self.task)['outputs']])
        emit_step(self.store, self.task, 'pivot_table', info)
        step = [e for e in self.store.events(self.task) if e['kind'] == 'step'][-1]
        self.assertEqual(step['data']['label'], STEP_LABELS['pivot_table'])
        self.assertEqual(step['data']['label'], '生成透视表')

    def test_rejects_multiple_inputs(self):
        fid = self.imported('数据.csv', '物料,数量\n螺丝,10\n')
        with self.assertRaisesRegex(ValueError, '需要一个明确输入'):
            Operation(kind='pivot_table', inputs=[{'file_id': fid}, {'file_id': fid}], params={'rows': ['物料'], 'values': [{'field': '数量'}]})

    def test_rejects_workbook_input(self):
        with self.assertRaisesRegex(ValueError, '透视表使用文件输入'):
            Operation(kind='pivot_table', inputs=[{'workbook_id': 'book1', 'version': 1, 'sheet_id': 'sheet1'}], params={'rows': ['物料'], 'values': [{'field': '数量'}]})


if __name__ == '__main__':
    unittest.main()
