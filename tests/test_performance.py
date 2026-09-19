import tempfile
import unittest
from pathlib import Path
from unittest import mock

import pandas as pd
from openpyxl import Workbook

from assistant.store import Store
from assistant.models import InputSelection
from assistant.tables import read, _read_csv_fast, _xlsx_has_formula


class CsvFastPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "app")
        self.task = self.store.create()["id"]

    def _csv(self, name, text, encoding="utf-8-sig"):
        path = self.root / name
        path.write_text(text, encoding=encoding)
        return path

    def test_fast_equals_pandas(self):
        path = self._csv("clean.csv", "code,qty\n001,1\n002,2\nNA,\n")
        with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 0):
            fast = _read_csv_fast(path)
        slow = pd.read_csv(path, header=None, dtype=str, keep_default_na=False, encoding="utf-8-sig", skip_blank_lines=False)
        self.assertIsNotNone(fast)
        self.assertTrue(fast.astype(object).equals(slow.astype(object)))
        self.assertEqual(fast.iloc[1, 0], "001")  # leading zero preserved
        self.assertEqual(fast.iloc[3, 1], "")  # empty field -> ''

    def test_small_file_falls_back(self):
        self.assertIsNone(_read_csv_fast(self._csv("small.csv", "a,b\n1,2\n")))

    def test_blank_line_falls_back(self):
        path = self._csv("blank.csv", "a,b\n1,2\n\n3,4\n")
        with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 0):
            self.assertIsNone(_read_csv_fast(path))

    def test_gb18030_falls_back(self):
        path = self._csv("gb.csv", "名称,数量\n北京,3\n", encoding="gb18030")
        with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 0):
            self.assertIsNone(_read_csv_fast(path))

    def test_read_fast_and_slow_match(self):
        path = self._csv("src.csv", "code,qty\nA,1\nB,2\nC,3\n")
        file = self.store.import_file(self.task, path)
        selection = InputSelection(file_id=file["id"], header_row=1, sheet=0)
        with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 0):
            fast = read(self.store, self.task, selection)
        with mock.patch("assistant.tables.CSV_FAST_MIN_BYTES", 10**12):
            slow = read(self.store, self.task, selection)
        self.assertEqual(fast.code.tolist(), ["A", "B", "C"])
        self.assertEqual(fast["__source_row"].tolist(), [2, 3, 4])
        self.assertEqual(fast["__source_row"].tolist(), slow["__source_row"].tolist())
        self.assertEqual(fast["__source_file"].tolist(), slow["__source_file"].tolist())


class XlsxFastPathTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "app")
        self.task = self.store.create()["id"]

    def _xlsx(self, name, rows, number_formats=None):
        path = self.root / name
        book = Workbook()
        for row in rows:
            book.active.append(row)
        for coord, fmt in (number_formats or {}).items():
            book.active[coord].number_format = fmt
        book.save(path)
        book.close()
        return path

    def _import(self, path):
        file = self.store.import_file(self.task, path)
        return InputSelection(file_id=file["id"], header_row=1, sheet=0)

    def test_has_formula_detection(self):
        plain = self._xlsx("plain.xlsx", [["a", "b"], [1, 2]])
        formula = self._xlsx("formula.xlsx", [["a"], [1], ["=A2+1"]])
        self.assertFalse(_xlsx_has_formula(plain))
        self.assertTrue(_xlsx_has_formula(formula))

    def test_value_only_read_keeps_zfill(self):
        path = self._xlsx("v.xlsx", [["code", "qty"], [12, 2]], number_formats={"A2": "00000"})
        frame = read(self.store, self.task, self._import(path))
        self.assertEqual(frame.code.iloc[0], "00012")
        self.assertEqual(frame.qty.iloc[0], 2)

    def test_single_pass_equals_double_pass(self):
        path = self._xlsx("v.xlsx", [["code", "qty"], ["NA", 1], ["00012", 2]])
        selection = self._import(path)
        single = read(self.store, self.task, selection)
        with mock.patch("assistant.tables._xlsx_has_formula", return_value=True):
            double = read(self.store, self.task, selection)
        self.assertEqual(single.code.tolist(), double.code.tolist())
        self.assertEqual(single["__source_row"].tolist(), double["__source_row"].tolist())

    def test_formula_without_cache_still_raises(self):
        path = self._xlsx("f.xlsx", [["code"], [1], ["=A2+1"]])
        with self.assertRaisesRegex(ValueError, "缓存"):
            read(self.store, self.task, self._import(path))


if __name__ == "__main__":
    unittest.main()
