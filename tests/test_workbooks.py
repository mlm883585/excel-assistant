import copy
from datetime import datetime
import tempfile
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.datetime import MAC_EPOCH
import xlsxwriter

from assistant.store import Store
from assistant.models import Operation, InputSelection
from assistant.tables import run_operation, read
from assistant.workbook_model import blank_workbook, blank_sheet, validate_snapshot, MAX_CELLS
from assistant.workbook_excel import import_workbook, export_workbook
from assistant.workbook_operations import create_table, edit_workbook
from assistant.workbooks import Workbooks


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.store = Store(self.root/'state'); self.task = self.store.create()['id']
        self.books = Workbooks(self.store)

    def imported(self, name, text):
        path = self.root/name; path.write_text(text, encoding='utf-8-sig')
        return self.store.import_file(self.task, path)['id']

    def test_concurrent_workbook_database_initialization(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier

        store = Store(self.root / 'fresh-state')
        task = store.create()['id']
        barrier = Barrier(6)

        def initialize(_):
            barrier.wait(timeout=10)
            return Workbooks(store).list(task)

        with ThreadPoolExecutor(max_workers=6) as pool:
            self.assertEqual(list(pool.map(initialize, range(6))), [[]] * 6)
        with store.connect() as db:
            names = [row[1] for row in db.execute('PRAGMA table_info(output_reviews)')]
            self.assertIn('run_id', names)
            self.assertIn('state', names)

    def test_date_serial_uses_relationship_and_gradient_is_readonly(self):
        from zipfile import ZipFile
        from openpyxl.styles import GradientFill
        source = self.root/'physical-order.xlsx'
        with xlsxwriter.Workbook(source) as book:
            first = book.add_worksheet('先表'); second = book.add_worksheet('后表')
            fmt = book.add_format({'num_format':'yyyy-mm-dd'})
            first.write_number(0,0,60,fmt); second.write_number(0,0,61,fmt)
        rewritten = self.root/'relationships.xlsx'
        with ZipFile(source) as old, ZipFile(rewritten,'w') as new:
            for item in old.infolist():
                data = old.read(item.filename)
                if item.filename == 'xl/_rels/workbook.xml.rels':
                    data = data.replace(b'worksheets/sheet1.xml',b'worksheets/sheet77.xml')
                name = 'xl/worksheets/sheet77.xml' if item.filename=='xl/worksheets/sheet1.xml' else item.filename
                new.writestr(name,data)
        snapshot, issues = import_workbook(rewritten)
        self.assertEqual(issues,[])
        self.assertEqual(snapshot['sheets'][0]['cells']['0,0']['value'],60)
        self.assertEqual(snapshot['sheets'][1]['cells']['0,0']['value'],61)
        gradient = self.root/'gradient.xlsx'
        b=Workbook();b.active['A1']='特殊填充';b.active['A1'].fill=GradientFill(stop=('FFFFFF','000000'));b.save(gradient);b.close()
        _, issues = import_workbook(gradient)
        self.assertTrue(any('特殊填充' in issue for issue in issues))

    def test_history_view_is_readonly_without_reverting_current(self):
        initial=self.books.open(self.task);snapshot=initial['snapshot']
        snapshot['sheets'][0]['cells']['0,0']={'value':'新内容','type':'string'}
        self.books.save(self.task,snapshot,1)
        old=self.books.open(self.task,snapshot['id'],revision=1)
        self.assertTrue(old['historical'] and old['readonly'])
        self.assertEqual(old['snapshot']['sheets'][0]['cells'],{})
        self.assertEqual(self.books.load(self.task,snapshot['id'])['revision'],2)
        restored=self.books.restore(self.task,snapshot['id'],1,2)
        self.assertEqual(restored['revision'],3)
        self.assertEqual(len(self.books.versions(self.task,snapshot['id'])),3)

    def test_rich_text_and_grouping_are_not_silently_discarded(self):
        source=self.root/'rich.xlsx'
        with xlsxwriter.Workbook(source) as book:
            sheet=book.add_worksheet()
            sheet.write_rich_string(0,0,'普通',book.add_format({'bold':True}),'加粗')
            sheet.set_row(1,None,None,{'level':1})
        _, issues=import_workbook(source)
        self.assertTrue(any('分段文字' in issue for issue in issues))
        self.assertTrue(any('分组' in issue for issue in issues))

    def test_csv_text_and_common_xlsx_roundtrip(self):
        fid = self.imported('编号.csv', '编号,内容\n00123,=SUM(A1:A2)\nNA,0\n')
        result = self.books.open(self.task, file_id=fid)
        s = result['snapshot']['sheets'][0]
        self.assertEqual(s['cells']['1,0']['value'], '00123')
        self.assertNotIn('formula', s['cells']['1,1'])
        target = self.root/'csv.xlsx'
        export_workbook(result['snapshot'], target)
        book = load_workbook(target); self.addCleanup(book.close)
        self.assertEqual(book.active['A2'].value, '00123')
        self.assertEqual(book.active['B2'].data_type, 's')

    def test_styles_structure_dates_and_epoch(self):
        for epoch1904 in (False, True):
            source = self.root/f'epoch-{epoch1904}.xlsx'
            b = Workbook()
            if epoch1904: b.epoch = MAC_EPOCH
            s = b.active; s.title = '业务'; b.create_sheet('汇总')
            s['A1'] = '物料'; s['A2'] = '0007'; s['B2'] = datetime(2024, 2, 29, 12)
            s['B2'].number_format = 'yyyy-mm-dd hh:mm'
            s['A1'].font = Font(name='宋体', size=14, bold=True, italic=True, color='112233')
            s['A1'].fill = PatternFill('solid', fgColor='AABBCC')
            s['A1'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            s['A1'].border = Border(bottom=Side(style='thin', color='445566'))
            s.merge_cells('A4:C4'); s['A4'] = '合并'; s.freeze_panes = 'B2'
            s.row_dimensions[2].height = 30; s.column_dimensions['C'].hidden = True
            s.column_dimensions['D'].font = Font(bold=True)
            b.save(source); b.close()
            data, issues = import_workbook(source)
            self.assertEqual(issues, [])
            target = self.root/f'out-{epoch1904}.xlsx'; export_workbook(data, target)
            check = load_workbook(target)
            try:
                self.assertEqual(check.sheetnames, ['业务', '汇总'])
                self.assertEqual(check.epoch.year, 1904 if epoch1904 else 1899)
                self.assertEqual(check['业务']['B2'].value, datetime(2024, 2, 29, 12))
                self.assertEqual(check['业务']['A2'].value, '0007')
                self.assertEqual(check['业务']['A1'].font.name, '宋体')
                self.assertTrue(check['业务']['A1'].font.bold)
                self.assertEqual(check['业务']['A1'].fill.fgColor.rgb[-6:], 'AABBCC')
                self.assertEqual(check['业务'].freeze_panes, 'B2')
                self.assertEqual(str(next(iter(check['业务'].merged_cells.ranges))), 'A4:C4')
                self.assertTrue(check['业务'].column_dimensions['C'].hidden)
                self.assertEqual(check['业务'].row_dimensions[2].height, 30)
            finally: check.close()

    def test_formula_cache_no_fabricated_zero(self):
        path = self.root/'formula.xlsx'
        with xlsxwriter.Workbook(path) as b:
            s = b.add_worksheet('明细'); s.write(0, 0, 3); s.write_formula(0, 1, '=A1*7', None, 21)
        data, issues = import_workbook(path); self.assertFalse(issues)
        self.assertEqual(data['sheets'][0]['cells']['0,1']['value'], 21)
        export_workbook(data, self.root/'formula-out.xlsx')
        check = load_workbook(self.root/'formula-out.xlsx', data_only=True)
        self.assertEqual(check.active['B1'].value, 21); check.close()
        cell = data['sheets'][0]['cells']['0,1']; cell['result_state']='pending'; cell['value']=None
        with self.assertRaisesRegex(ValueError, '重算'): export_workbook(data, self.root/'bad.xlsx')
        cell.update(result_state='error', value='#DIV/0!', type='error')
        with self.assertRaisesRegex(ValueError, '重算'): export_workbook(data, self.root/'bad.xlsx')

    def test_unsupported_features_are_readonly_and_extract_is_explicit(self):
        from openpyxl.chart import BarChart, Reference
        b = Workbook(); b.active.append(['值']); b.active.append([2])
        chart = BarChart(); chart.add_data(Reference(b.active, min_col=1,min_row=1,max_row=2)); b.active.add_chart(chart,'D1')
        path = self.root/'chart.xlsx'; b.save(path); b.close()
        fid = self.store.import_file(self.task,path)['id']
        record = self.books.open(self.task,file_id=fid)
        self.assertTrue(record['readonly']); self.assertTrue(record['limitations'])
        with self.assertRaises(ValueError): self.books.save(self.task,record['snapshot'],1)
        pure = self.books.open(self.task,file_id=fid,pure_data=True)
        self.assertFalse(pure['readonly']); self.assertNotEqual(pure['snapshot']['id'],record['snapshot']['id'])

    def test_version_conflict_restore_atomic_failure_and_task_isolation(self):
        record = self.books.open(self.task); data=record['snapshot']; data['sheets'][0]['cells']['0,0']={'value':'更改','type':'string'}
        self.assertEqual(self.books.save(self.task,data,1)['revision'],2)
        with self.assertRaisesRegex(ValueError,'版本冲突'): self.books.save(self.task,data,1)
        self.assertEqual(self.books.load(self.task,data['id'])['revision'],2)
        restored=self.books.restore(self.task,data['id'],1,2)
        self.assertEqual(restored['revision'],3); self.assertEqual(restored['snapshot']['sheets'][0]['cells'],{})
        self.assertEqual(len(self.books.versions(self.task,data['id'])),3)
        with patch('assistant.workbooks.packed',side_effect=OSError('disk full')):
            with self.assertRaises(OSError): self.books.save(self.task,data,3)
        self.assertEqual(self.books.load(self.task,data['id'])['revision'],3)
        with self.assertRaises(ValueError): self.books.load(self.store.create()['id'],data['id'])

    def test_capacity_union_and_row_styles(self):
        data=blank_workbook(); s=data['sheets'][0]; s['row_count']=10000
        s['cells']={f'{i//20},{i%20}':{'value':i,'style':{'bold':True}} for i in range(MAX_CELLS)}
        s['columns']['0']={'style':{'bg_color':'#112233'}}
        self.assertEqual(validate_snapshot(data),MAX_CELLS)
        s['cells']['0,20']={'style':{'italic':True}}
        with self.assertRaisesRegex(ValueError,'200000'):validate_snapshot(data)

    def test_calculate_classify_errors_and_rules_validation(self):
        fid=self.imported('业务.csv','编号,数量,单价\n001,2,3.5\n002,,4\n003,非法,5\n004,5,0\n')
        op={'kind':'calculate','inputs':[{'file_id':fid}],'params':{'columns':[{'name':'金额','expression':{'op':'multiply','args':[{'field':'数量'},{'field':'单价'}]}}]}}
        output=run_operation(self.store,self.task,op)
        data=read(self.store,self.task,InputSelection(file_id=output['id']))
        self.assertEqual(len(data),4); self.assertEqual(data['金额'].iloc[0],7)
        self.assertEqual(output['statistics']['issues'],2)
        review=self.books.review(self.task,output['id']); self.assertGreater(review['total'],0)
        with self.assertRaises(ValueError):self.books.changes(self.task,output['id'],limit=201)
        bad=copy.deepcopy(op); bad['params']['columns'][0]['expression']={'code':'import os'}
        with self.assertRaises(ValueError):run_operation(self.store,self.task,bad)
        bad=copy.deepcopy(op); bad['params']['columns'][0]['name']='数量'
        with self.assertRaises(ValueError):run_operation(self.store,self.task,bad)
        cls={'kind':'classify','inputs':[{'file_id':fid}],'params':{'columns':[{'name':'等级','rules':[{'when':{'field':'数量','operator':'ge','value':3},'label':'高'}],'default':'低'}]}}
        info=run_operation(self.store,self.task,cls); rows=read(self.store,self.task,InputSelection(file_id=info['id']))
        self.assertEqual(rows['等级'].tolist(),['低','','','高'])

    def test_create_table_zero_input_recipe_and_no_invented_data(self):
        op={'kind':'create_table','inputs':[],'params':{'columns':['编号','数量']}}
        info=run_operation(self.store,self.task,op)
        adopted=self.books.apply(self.task,info['id']); sheet=adopted['snapshot']['sheets'][0]
        self.assertEqual(sheet['row_count'],21); self.assertTrue(all(c.get('value') is None for k,c in sheet['cells'].items() if not k.startswith('0,')))
        task=self.store.get(self.task); task.update(status='succeeded',plan={'steps':[op],'questions':[]}); self.store.save(task)
        recipe=self.store.save_recipe(self.task,'空白表'); self.assertEqual(recipe['slots'],[])
        self.assertEqual(self.store.recipe_plan(recipe['id'],self.task)['steps'][0]['inputs'],[])

    def test_workbook_selection_candidate_stale_and_cancel(self):
        record=self.books.open(self.task); data=create_table({'columns':['编号','数量','单价'],'rows':[['01',2,5]],'blank_rows':0})
        data['id']=record['snapshot']['id']; self.books.save(self.task,data,1)
        sid=data['sheets'][0]['id']
        selection={'workbook_id':data['id'],'version':2,'sheet_id':sid,'range':{'r0':0,'c0':0,'r1':1,'c1':2}}
        frame=read(self.store,self.task,InputSelection(**selection)); self.assertEqual(frame['__source_row'].tolist(),[2])
        op={'kind':'edit_workbook','inputs':[selection],'params':{'edits':[{'kind':'set_values','range':{'r0':1,'c0':1,'r1':1,'c1':1},'values':[[9]]}]}}
        info=run_operation(self.store,self.task,op)
        self.assertEqual(self.books.load(self.task,data['id'])['snapshot']['sheets'][0]['cells']['1,1']['value'],2)
        self.books.save(self.task,data,2)
        with self.assertRaisesRegex(ValueError,'版本冲突'):self.books.apply(self.task,info['id'])
        op['inputs'][0]['version']=3
        info2=run_operation(self.store,self.task,op)
        task=self.store.get(self.task); task['status']='cancelled';self.store.save(task)
        with self.assertRaisesRegex(ValueError,'取消'):self.books.apply(self.task,info2['id'])
        task['status']='succeeded';self.store.save(task)
        applied=self.books.apply(self.task,info2['id']);self.assertEqual(applied['revision'],4)
        self.assertEqual(applied['snapshot']['sheets'][0]['cells']['1,1']['value'],9)

    def test_export_exclusive_revision_and_busy(self):
        record=self.books.open(self.task); bid=record['snapshot']['id']; target=self.root/'result.xlsx'
        self.assertTrue(self.books.export(self.task,bid,1,target)['exported'])
        with self.assertRaisesRegex(ValueError,'新文件'):self.books.export(self.task,bid,1,target)
        self.books.save(self.task,record['snapshot'],1)
        with self.assertRaisesRegex(ValueError,'失效'):self.books.export(self.task,bid,1,self.root/'stale.xlsx')
        task=self.store.get(self.task);task['status']='running';self.store.save(task)
        with self.assertRaisesRegex(ValueError,'只读'):self.books.save(self.task,record['snapshot'],2)

    def test_failed_export_removes_partial_copy_without_marking_success(self):
        import shutil

        record = self.books.open(self.task)
        target = self.root / 'failed-copy.xlsx'
        copy_file = shutil.copyfileobj

        def fail_after_write(source, destination, *args):
            # ZipFile also uses copyfileobj internally. Inject the failure only
            # into the application's final exclusive destination copy.
            if getattr(destination, 'name', None) != str(target):
                return copy_file(source, destination, *args)
            destination.write(source.read(32))
            raise OSError('synthetic disk full')

        with patch('shutil.copyfileobj', side_effect=fail_after_write):
            with self.assertRaisesRegex(OSError, 'disk full'):
                self.books.export(self.task, record['snapshot']['id'], 1, target)
        self.assertFalse(target.exists())
        self.assertFalse(list(self.root.glob('.workbook-*.xlsx')))
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT count(*) FROM workbook_exports').fetchone()[0], 0)
        self.assertEqual(self.books.load(self.task, record['snapshot']['id'])['revision'], 1)

    def test_malformed_and_arbitrary_commands_rejected(self):
        with self.assertRaises(ValueError):Operation.model_validate({'kind':'calculate','inputs':[]})
        with self.assertRaises(ValueError):InputSelection(file_id='x',workbook_id='y',version=1,sheet_id='s')
        book=blank_workbook()
        with self.assertRaises(ValueError):edit_workbook(book,book['sheets'][0]['id'],{'edits':[{'kind':'run_command','command':'anything'}]})
        with self.assertRaises(ValueError):edit_workbook(book,book['sheets'][0]['id'],{'edits':[{'kind':'set_style','range':{'r0':0,'c0':0,'r1':0,'c1':0},'style':{'unknown':1}}]})

    def test_cancelled_candidate_does_not_revive_on_later_success(self):
        task=self.store.get(self.task);task.update(status='running',run_id='first');self.store.save(task)
        info=run_operation(self.store,self.task,{'kind':'create_table','params':{'columns':['字段']}})
        self.assertEqual(self.books.review(self.task,info['id'])['state'],'pending')
        self.store.finish_reviews(self.task,False)
        task=self.store.get(self.task);task.update(status='succeeded',run_id='second');self.store.save(task)
        with self.assertRaisesRegex(ValueError,'取消'):self.books.apply(self.task,info['id'])

    def test_numeric_date_serial_60_and_time_1904_survive(self):
        from zipfile import ZipFile
        for mac in (False, True):
            path=self.root/f'dates-{mac}.xlsx'
            with xlsxwriter.Workbook(path,{'date_1904':mac}) as book:
                sheet=book.add_worksheet();fmt=book.add_format({'num_format':'yyyy-mm-dd'})
                sheet.write_number(0,0,60,fmt);sheet.write_number(1,0,0,fmt)
                sheet.write_number(2,0,.5,book.add_format({'num_format':'hh:mm'}))
            data,issues=import_workbook(path);self.assertFalse(issues)
            self.assertEqual(data['sheets'][0]['cells']['0,0']['value'],60+1462 if mac else 60)
            target=self.root/f'dates-out-{mac}.xlsx';export_workbook(data,target)
            with ZipFile(target) as archive:
                xml=archive.read('xl/worksheets/sheet1.xml').decode()
                self.assertIn('<v>60</v>',xml);self.assertIn('<v>0.5</v>',xml)

    def test_selected_edit_range_is_enforced(self):
        record=self.books.open(self.task);book=record['snapshot'];sid=book['sheets'][0]['id']
        op={'kind':'edit_workbook','inputs':[{'workbook_id':book['id'],'version':1,'sheet_id':sid,'range':{'r0':0,'c0':0,'r1':0,'c1':0}}],'params':{'edits':[{'kind':'set_values','range':{'r0':1,'c0':0,'r1':1,'c1':0},'values':[[7]]}]}}
        with self.assertRaisesRegex(ValueError,'超出用户选区'):run_operation(self.store,self.task,op)
