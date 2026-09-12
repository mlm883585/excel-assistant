import asyncio
import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from assistant.store import Store
from assistant.tables import read, run_operation, preview
from assistant.models import InputSelection
from assistant.agent import AgentRunner


class BusinessTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "app")
        self.task = self.store.create()["id"]

    def tearDown(self):
        self.temp.cleanup()

    def file(self, rows, header=1):
        path = self.root / f"input{len(self.store.get(self.task)['files'])}.xlsx"
        book = Workbook()
        for row in rows:
            book.active.append(row)
        book.save(path)
        book.close()
        file = self.store.import_file(self.task, path)
        return {"file_id": file["id"], "header_row": header, "sheet": 0}

    def execute(self, kind, inputs, **params):
        return run_operation(self.store, self.task, {"kind": kind, "inputs": inputs, "params": params})

    def output(self, result):
        return preview(self.store, self.task, {"file_id":result["id"]})

    def test_na_leading_zero_and_real_source_row(self):
        item = self.file([["title", ""], ["说明", ""], ["code", "qty"], ["NA", 1], ["00012", 2]], header=3)
        frame = read(self.store, self.task, InputSelection(**item))
        self.assertEqual(frame.code.tolist(), ["NA", "00012"])
        self.assertEqual(frame.__getitem__("__source_row").tolist(), [4, 5])

    def test_formatted_code_and_missing_formula_cache(self):
        item = self.file([["code", "qty"], [12, 2]])
        path, _ = self.store.resolve_file(self.task, item["file_id"])
        book = load_workbook(path)
        book.active['A2'].number_format = '00000'
        book.save(path)
        book.close()
        self.assertEqual(read(self.store,self.task,InputSelection(**item)).code.iloc[0], '00012')
        book = load_workbook(path)
        book.active['B2'] = '=1+1'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValueError, '缓存'):
            read(self.store, self.task, InputSelection(**item))

    def test_append_preview_and_formula_like_text(self):
        a = self.file([["code"], ["=ABC"]])
        # Excel strings beginning with '=' must explicitly be stored as text.
        path, _ = self.store.resolve_file(self.task, a['file_id'])
        book = load_workbook(path); book.active['A2'].data_type='s'; book.save(path); book.close()
        b = self.file([["code"], ["NA"]])
        result = self.execute('append',[a,b])
        self.assertEqual(self.output(result)['total'],2)
        path, _ = self.store.resolve_file(self.task,result['id'])
        book=load_workbook(path); self.assertEqual(book.active['A2'].data_type,'s'); book.close()

    def test_join_duplicates_rejected(self):
        a=self.file([["code","qty"],["A",1],["A",2]])
        b=self.file([["code","stock"],["A",3],["A",4]])
        with self.assertRaisesRegex(Exception,'不唯一'):
            self.execute('join',[a,b],keys=['code'])
        with self.assertRaisesRegex(ValueError,'多对多'):
            self.execute('join',[a,b],keys=['code'],validate='many_to_many')

    def test_join_unmatched_is_exported(self):
        a=self.file([["code","qty"],["A",1],["B",2]])
        b=self.file([["code","stock"],["A",3]])
        result=self.execute('join',[a,b],keys=['code'])
        path,_=self.store.resolve_file(self.task,result['id'])
        book=load_workbook(path); self.assertEqual(book['左表未匹配']['A2'].value,'B'); book.close()

    def test_clean_error_row_matches_source(self):
        a=self.file([["title"],["qty"],["bad"]],header=2)
        result=self.execute('clean',[a],config={'schema_version':1,'columns':{'qty':{'type':'integer'}}})
        self.assertEqual(result['issues'][0]['行号'],3)

    def test_compare_added_removed_and_changes(self):
        a=self.file([["code","qty"],["A",1],["B",2]])
        b=self.file([["code","qty"],["A",3],["C",2]])
        result=self.execute('compare',[a,b],keys=['code'])
        self.assertEqual(self.output(result)['total'],1)
        self.assertEqual(self.output(result)['rows'][0]['右值'],3)

    def test_group_and_matrix_roundtrip(self):
        a=self.file([["code","month","qty"],["A","Jan",2],["A","Feb",3]])
        result=self.execute('group',[a],keys=['code'],columns=['qty'])
        self.assertEqual(self.output(result)['rows'][0]['qty'],5)
        result=self.execute('pivot',[a],keys=['code'],column='month',value='qty')
        self.assertEqual(self.output(result)['rows'][0]['Jan'],2)
        result=self.execute('melt',[{'file_id':result['id']}],keys=['code'],columns=['Jan','Feb'])
        self.assertEqual(self.output(result)['total'],2)

    def test_template_preserves_formula_and_source(self):
        a=self.file([["code","qty"],["A",4]])
        b=self.file([["code","qty","double"],["",0,"=B2*2"]])
        path,_=self.store.resolve_file(self.task,b['file_id']); original=path.read_bytes()
        result=self.execute('template',[a,b],sheet='Sheet',mapping={'code':'A','qty':'B'})
        output,_=self.store.resolve_file(self.task,result['id'])
        book=load_workbook(output); self.assertEqual(book.active['B2'].value,4); self.assertEqual(book.active['C2'].value,'=B2*2'); book.close()
        self.assertEqual(path.read_bytes(),original)
        with self.assertRaisesRegex(ValueError,'已有公式'):
            self.execute('template',[a,b],mapping={'qty':'C'})

    def test_path_and_pagination_boundaries(self):
        a=self.file([["code"],["A"]])
        with self.assertRaises(ValueError): self.store.resolve_file(self.task,'../../secret')
        with self.assertRaises(ValueError): self.store.get('../outside')
        with self.assertRaises(ValueError): preview(self.store,self.task,a,limit=201)

    def test_recipe_rebinds_files(self):
        a=self.file([["code"],["A"]])
        record=self.store.get(self.task); record.update(status='succeeded',plan={'steps':[{'kind':'append','inputs':[a],'params':{}}],'questions':[]}); self.store.save(record)
        recipe=self.store.save_recipe(self.task,'monthly')
        new=self.store.create()['id']
        source=self.root/'new.csv'; source.write_text('code\nB\n')
        replacement=self.store.import_file(new,source)
        plan=self.store.recipe_plan(recipe['id'],new)
        self.assertEqual(plan['steps'][0]['inputs'][0]['file_id'],replacement['id'])

    def test_hundred_thousand_rows_complete(self):
        path=self.root/'large.csv'
        path.write_text('code,qty\n' + ''.join(f'{i:06d},{i}\n' for i in range(100000)))
        item=self.store.import_file(self.task,path)
        result=self.execute('append',[{'file_id':item['id']}])
        page=preview(self.store,self.task,{'file_id':result['id']},offset=99999,limit=1)
        self.assertEqual(page['total'],100000)
        self.assertEqual(page['rows'][0]['code'],'099999')

    def test_agent_permissions_deny_shell_and_other_servers(self):
        runner=AgentRunner(self.store,self.task,{})
        for name in ['run_shell_command','read_file','mcp__evil__datacraft_execute']:
            self.assertEqual(asyncio.run(runner.permission(name,{},{}))['behavior'],'deny')
        self.assertEqual(asyncio.run(runner.permission('mcp__datacraft__datacraft_preview',{},{}))['behavior'],'allow')

    def test_model_claim_is_not_success(self):
        class Stream:
            async def __aenter__(self): return self
            async def __aexit__(self,*args): pass
            def __aiter__(self): return self.generate()
            async def generate(self): yield {'type':'result','result':'done'}
        runner=AgentRunner(self.store,self.task,{'base_url':'http://localhost:9999/v1','model':'test','cli':__file__},lambda *args:Stream())
        with self.assertRaisesRegex(ValueError,'尚未生成'):
            asyncio.run(runner.run('test'))


if __name__=='__main__': unittest.main()
