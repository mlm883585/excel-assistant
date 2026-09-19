"""Generate-only native Excel pivot table writer.

openpyxl (3.1.5) exposes only read-side pivot classes and XlsxWriter has no pivot
support, so a real Excel pivot table is produced by writing the source data with
XlsxWriter and then injecting the three pivot OOXML parts (pivotTable,
pivotCacheDefinition, pivotCacheRecords) into the zip.

The result is a one-shot output: re-importing it still lands in the read-only path
(``xl/pivotTables/`` preflight in :mod:`assistant.workbook_excel`). Nothing here is
round-tripped back into the editor.
"""

import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

CT_WORKSHEET = "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
CT_PIVOT_TABLE = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"
CT_PIVOT_CACHE_DEF = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheDefinition+xml"
CT_PIVOT_CACHE_REC = "application/vnd.openxmlformats-officedocument.spreadsheetml.pivotCacheRecords+xml"

REL_PIVOT_TABLE = f"{REL_NS}/pivotTable"
REL_PIVOT_CACHE_DEF = f"{REL_NS}/pivotCacheDefinition"
REL_PIVOT_CACHE_REC = f"{REL_NS}/pivotCacheRecords"

AGGREGATE_SUBTOTAL = {"sum": "sum", "count": "count", "mean": "average", "min": "min", "max": "max"}
AGGREGATE_CAPTION = {"sum": "求和项", "count": "计数项", "mean": "平均值项", "min": "最小值项", "max": "最大值项"}


def _esc(value):
    return escape(str(value), {'"': "&quot;"})


def _blank(value):
    return value is None or (isinstance(value, float) and pd.isna(value))


def _coerce(value):
    number = float(value)
    return int(number) if number.is_integer() else number


def _key(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return ("n", value)
    return ("s", value)


def _fmt_number(value):
    number = float(value)
    return str(int(number)) if number.is_integer() else repr(number)


def _column_profile(series, force_number=False):
    """Return (is_number, coerced_series, index_map, distinct_values) for a column."""
    if force_number:
        coerced = pd.to_numeric(series, errors="raise")
        number = True
    else:
        non_blank = [v for v in series.tolist() if not _blank(v)]
        number = bool(non_blank) and all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in non_blank)
        if number:
            coerced = pd.Series([_coerce(v) if not _blank(v) else v for v in series.tolist()], index=series.index)
        else:
            coerced = pd.Series(["" if _blank(v) else str(v) for v in series.tolist()], index=series.index)
    order, index, values = [], {}, []
    for value in coerced.tolist():
        key = _key(value)
        if key not in index:
            index[key] = len(order)
            order.append(key)
            values.append(value)
    return number, coerced, index, values


def _shared_items_xml(values, number):
    attrs = [f'count="{len(values)}"']
    if number:
        integral = all(float(v) == int(float(v)) for v in values)
        attrs += ['containsSemiMixedTypes="0"', 'containsString="0"', 'containsNumber="1"',
                  f'containsInteger="{"1" if integral else "0"}"',
                  f'minValue="{_fmt_number(min(float(v) for v in values))}"',
                  f'maxValue="{_fmt_number(max(float(v) for v in values))}"']
        body = "".join(f'<n v="{_fmt_number(v)}"/>' for v in values)
    else:
        attrs += ['containsSemiMixedTypes="0"', 'containsString="1"', 'containsNumber="0"', 'containsInteger="0"']
        body = "".join(f'<s v="{_esc(v)}"/>' for v in values)
    return f'<sharedItems {" ".join(attrs)}>{body}</sharedItems>'


def _cache_field_xml(name, values, number):
    return f'<cacheField name="{_esc(name)}" numFmtId="0">{_shared_items_xml(values, number)}</cacheField>'


def _cache_definition_xml(record_count, ref, sheet, fields):
    fields_xml = "".join(_cache_field_xml(name, values, number) for name, values, number in fields)
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheDefinition xmlns="{MAIN_NS}" xmlns:r="{REL_NS}" r:id="rId1" '
        f'recordCount="{record_count}" refreshOnLoad="1" saveData="1" invalid="0" '
        f'enableRefresh="1" createdVersion="8" refreshedVersion="8" minRefreshableVersion="3">'
        f'<cacheSource type="worksheet"><worksheetSource ref="{ref}" sheet="{_esc(sheet)}"/></cacheSource>'
        f'<cacheFields count="{len(fields)}">{fields_xml}</cacheFields>'
        f'</pivotCacheDefinition>'
    )


def _cache_records_xml(records):
    rows = "".join("<r>" + "".join(f'<x v="{i}"/>' for i in record) + "</r>" for record in records)
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotCacheRecords xmlns="{MAIN_NS}" xmlns:r="{REL_NS}" count="{len(records)}">{rows}</pivotCacheRecords>'
    )


def _pivot_field_xml(axis, items):
    body = "".join(f'<item x="{i}"/>' for i in items) + '<item t="default"/>'
    return f'<pivotField axis="{axis}" showAll="0"><items count="{len(items) + 1}">{body}</items></pivotField>'


def _pivot_table_xml(meta, pivot_field_parts, row_indices, col_index, value_index, location):
    rowf = f'<rowFields count="{len(row_indices)}">' + "".join(f'<field x="{i}"/>' for i in row_indices) + '</rowFields>'
    colf = f'<colFields count="1"><field x="{col_index}"/></colFields>' if col_index is not None else ''
    dataf = (f'<dataFields count="1"><dataField name="{_esc(meta["value_field"])}" fld="{value_index}" '
             f'subtotal="{AGGREGATE_SUBTOTAL[meta["aggregate"]]}" baseField="-1" baseItem="1048832"/></dataFields>')
    style = ('<pivotTableStyleInfo name="PivotStyleLight16" showRowHeaders="1" showColHeaders="1" '
             'showRowStripes="0" showColStripes="0" showLastColumn="1"/>')
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<pivotTableDefinition xmlns="{MAIN_NS}" xmlns:r="{REL_NS}" name="{_esc(meta["name"])}" '
        f'cacheId="1" dataCaption="{_esc(meta["data_caption"])}" dataOnRows="0" '
        f'createdVersion="8" updatedVersion="8" minRefreshableVersion="3" useAutoFormatting="1" itemPrintTitles="1">'
        f'{location}'
        f'<pivotFields count="{len(pivot_field_parts)}">{"".join(pivot_field_parts)}</pivotFields>'
        f'{rowf}{colf}{dataf}{style}'
        f'</pivotTableDefinition>'
    )


def _pivot_sheet_xml():
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<worksheet xmlns="{MAIN_NS}" xmlns:r="{REL_NS}"><sheetData/><pivotTableDefinition r:id="rId1"/></worksheet>'
    )


def _rels_xml(relationships):
    items = "".join(
        f'<Relationship Id="{rel_id}" Type="{rel_type}" Target="{target}"/>'
        for rel_id, rel_type, target in relationships
    )
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<Relationships xmlns="{PKG_REL_NS}">{items}</Relationships>'
    )


def _patch_content_types(xml):
    overrides = (
        f'<Override PartName="/xl/worksheets/sheet2.xml" ContentType="{CT_WORKSHEET}"/>'
        f'<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="{CT_PIVOT_TABLE}"/>'
        f'<Override PartName="/xl/pivotCache/pivotCacheDefinition1.xml" ContentType="{CT_PIVOT_CACHE_DEF}"/>'
        f'<Override PartName="/xl/pivotCache/pivotCacheRecords1.xml" ContentType="{CT_PIVOT_CACHE_REC}"/>'
    )
    return xml.replace("</Types>", overrides + "</Types>")


def _next_rid(rels_xml):
    ids = [int(m) for m in re.findall(r'Id="rId(\d+)"', rels_xml)]
    return max(ids, default=0) + 1


def _patch_workbook_rels(xml, sheet_rid, cache_rid):
    extra = (
        f'<Relationship Id="rId{sheet_rid}" Type="{REL_NS}/worksheet" Target="worksheets/sheet2.xml"/>'
        f'<Relationship Id="rId{cache_rid}" Type="{REL_PIVOT_CACHE_DEF}" Target="pivotCache/pivotCacheDefinition1.xml"/>'
    )
    return xml.replace("</Relationships>", extra + "</Relationships>")


def _patch_workbook(xml, sheet_name, sheet_id, sheet_rid, cache_rid):
    sheet = f'<sheet name="{_esc(sheet_name)}" sheetId="{sheet_id}" r:id="rId{sheet_rid}"/>'
    xml = xml.replace("</sheets>", sheet + "</sheets>")
    caches = f'<pivotCaches><pivotCache cacheId="1" r:id="rId{cache_rid}"/></pivotCaches>'
    return xml.replace("</workbook>", caches + "</workbook>")


def _inject_pivot(src, dst, parts, patch):
    with ZipFile(src) as source, ZipFile(dst, "w", ZIP_DEFLATED) as target:
        for name in source.namelist():
            content = source.read(name)
            if name in patch:
                content = patch[name](content.decode("utf-8")).encode("utf-8")
            target.writestr(name, content)
        for name, content in parts.items():
            target.writestr(name, content)


def build_pivot_xlsx(frame, spec, target):
    """Write a native Excel pivot table from ``frame`` into ``target`` (xlsx).

    ``spec`` follows the ``pivot_table`` operation contract::

        {"rows": ["物料"], "columns": ["地区"],
         "values": [{"field": "数量", "aggregate": "sum"}], "name": "透视表"}

    rows >= 1, columns <= 1, values == 1. The value field must be numeric.
    """
    target = Path(target)
    rows = spec.get("rows") or []
    columns = spec.get("columns") or []
    values = spec.get("values") or []
    if not rows:
        raise ValueError("透视表至少需要一个行字段")
    if len(columns) > 1:
        raise ValueError("透视表最多支持一个列字段")
    if len(values) != 1:
        raise ValueError("透视表需要一个数值字段")
    value_field = values[0].get("field")
    aggregate = values[0].get("aggregate", "sum")
    if aggregate not in AGGREGATE_SUBTOTAL:
        raise ValueError("请选择有效汇总规则")
    if value_field in rows or value_field in columns:
        raise ValueError("数值字段不能同时作为行或列字段")
    project = [*rows, *columns, value_field]
    missing = [c for c in project if c not in frame.columns]
    if missing:
        raise ValueError("请选择存在的字段")
    if frame.empty:
        raise ValueError("没有可透视的数据")

    profiles = {col: _column_profile(frame[col], force_number=(col == value_field)) for col in project}
    data = pd.DataFrame({col: profiles[col][1] for col in project})

    cache_fields = [(col, profiles[col][3], profiles[col][0]) for col in project]
    index_maps = [profiles[col][2] for col in project]
    records = []
    for row in data.itertuples(index=False, name=None):
        records.append([index_maps[i][_key(value)] for i, value in enumerate(row)])

    row_indices = list(range(len(rows)))
    col_index = len(rows) if columns else None
    value_index = len(project) - 1

    n_row_items = int(data[rows].drop_duplicates().shape[0])
    n_col_items = int(data[columns[0]].nunique()) if columns else 1
    header_rows = 1 + (1 if columns else 0)
    first_data_row = header_rows + 1
    first_data_col = len(rows) + 1
    total_cols = len(rows) + (n_col_items if columns else 1) + (1 if columns else 0)
    total_rows = header_rows + n_row_items + 1
    ref = f"A1:{get_column_letter(total_cols)}{total_rows}"
    location = f'<location ref="{ref}" firstHeaderRow="1" firstDataRow="{first_data_row}" firstDataCol="{first_data_col}"/>'

    pivot_field_parts = []
    for index, col in enumerate(project):
        values_in_field = profiles[col][3]
        if col == value_field:
            pivot_field_parts.append('<pivotField dataField="1" showAll="0"/>')
        elif index in row_indices:
            pivot_field_parts.append(_pivot_field_xml("axisRow", list(range(len(values_in_field)))))
        else:
            pivot_field_parts.append(_pivot_field_xml("axisCol", list(range(len(values_in_field)))))

    pivot_name = spec.get("name") or "数据透视表1"
    pivot_sheet = spec.get("name") or "透视表"
    meta = {
        "name": pivot_name, "value_field": value_field, "aggregate": aggregate,
        "data_caption": AGGREGATE_CAPTION[aggregate],
    }

    parts = {
        "xl/pivotTables/pivotTable1.xml": _pivot_table_xml(meta, pivot_field_parts, row_indices, col_index, value_index, location),
        "xl/pivotTables/_rels/pivotTable1.xml.rels": _rels_xml([("rId1", REL_PIVOT_CACHE_DEF, "../pivotCache/pivotCacheDefinition1.xml")]),
        "xl/pivotCache/pivotCacheDefinition1.xml": _cache_definition_xml(
            len(records), f"A1:{get_column_letter(len(project))}{len(data) + 1}", "数据", cache_fields),
        "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": _rels_xml([("rId1", REL_PIVOT_CACHE_REC, "pivotCacheRecords1.xml")]),
        "xl/pivotCache/pivotCacheRecords1.xml": _cache_records_xml(records),
        "xl/worksheets/sheet2.xml": _pivot_sheet_xml(),
        "xl/worksheets/_rels/sheet2.xml.rels": _rels_xml([("rId1", REL_PIVOT_TABLE, "../pivotTables/pivotTable1.xml")]),
    }

    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(f".{target.stem}.tmp.xlsx")
    built = target.with_name(f".{target.stem}.built.xlsx")
    try:
        with pd.ExcelWriter(temp, engine="xlsxwriter",
                            engine_kwargs={"options": {"strings_to_formulas": False, "strings_to_urls": False}}) as writer:
            data.to_excel(writer, index=False, sheet_name="数据")
            sheet = writer.sheets["数据"]
            sheet.freeze_panes(1, 0)
            if len(data.columns):
                sheet.autofilter(0, 0, len(data), len(data.columns) - 1)
                sheet.set_column(0, len(data.columns) - 1, 20)

        with ZipFile(temp) as source:
            rels = source.read("xl/_rels/workbook.xml.rels").decode("utf-8")
            workbook_xml = source.read("xl/workbook.xml").decode("utf-8")
        sheet_rid = _next_rid(rels)
        cache_rid = sheet_rid + 1
        sheet_id = max(int(m) for m in re.findall(r'sheetId="(\d+)"', workbook_xml)) + 1

        def _rels_patch(xml):
            return _patch_workbook_rels(xml, sheet_rid, cache_rid)

        def _workbook_patch(xml):
            return _patch_workbook(xml, pivot_sheet, sheet_id, sheet_rid, cache_rid)

        _inject_pivot(temp, built, parts, {
            "[Content_Types].xml": _patch_content_types,
            "xl/workbook.xml": _workbook_patch,
            "xl/_rels/workbook.xml.rels": _rels_patch,
        })
        check = load_workbook(built, read_only=True)
        check.close()
        built.replace(target)
    finally:
        temp.unlink(missing_ok=True)
        built.unlink(missing_ok=True)
