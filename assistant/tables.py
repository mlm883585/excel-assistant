import json
import re
import uuid
import time
from pathlib import Path

import pandas as pd
import python_calamine  # Explicit import ensures the xls reader is bundled.
from openpyxl import load_workbook

from data_toolkit.config import parse_config
from data_toolkit.cleaning import clean_dataframe
from data_toolkit.reconcile import merge_dataframes, compare_dataframes
from .models import InputSelection, Operation

SOURCE = ["__source_file", "__source_sheet", "__source_row"]

CHART_TYPES = {"column", "bar", "line", "area", "pie"}
CHART_AGGREGATES = {"sum", "count", "min", "max", "mean"}

STEP_LABELS = {"append": "合并文件", "join": "按字段关联", "clean": "清理文本", "compare": "对账差异",
               "group": "分组汇总", "melt": "宽表转长表", "pivot": "长表转矩阵", "template": "填写模板",
               "recalculate": "Excel 重算", "calculate": "新增计算列", "classify": "条件分级",
               "create_table": "新建工作簿", "edit_workbook": "编辑工作簿", "report": "数据质量报告",
               "chart": "生成图表", "pivot_table": "生成透视表", "sql": "SQL 关联"}


def emit_step(store, task_id, kind, info):
    """Record a conversation-visible step card from an operation's output info."""
    stats = info.get("statistics") or {}
    store.event(task_id, "step", {"kind": kind, "label": STEP_LABELS.get(kind, kind),
                                  "output_id": info.get("id"), "output_name": info.get("name"),
                                  "rows_in": stats.get("input_rows"), "rows_out": stats.get("output_rows"),
                                  "issues": stats.get("issues")})


def records(frame):
    return json.loads(frame.to_json(orient="records", date_format="iso", force_ascii=False))


def overview(path):
    if path.suffix == ".csv":
        return {"sheets": ["CSV"], "note": "请选择实际表头行；不自动猜测业务标题"}
    engine = "openpyxl" if path.suffix == ".xlsx" else "calamine"
    with pd.ExcelFile(path, engine=engine) as book:
        return {"sheets": book.sheet_names, "note": "请选择工作表及实际表头行"}


def read(store, task, selection: InputSelection):
    header = selection.header_row - 1
    row_offset = 0
    if selection.workbook_id:
        from .workbooks import Workbooks
        rows, info = Workbooks(store).input(task, selection)
        raw, sheet_name, row_offset = pd.DataFrame(rows), info['sheet'], info['row_offset']
        path = None
    else:
        path, info = store.resolve_file(task, selection.file_id)
    if path is None:
        pass
    elif info.get('kind') == 'workbook_candidate':
        from .workbook_model import validate_snapshot
        snapshot = json.loads(path.read_text(encoding='utf-8')); validate_snapshot(snapshot)
        candidate = snapshot['sheets'][selection.sheet] if isinstance(selection.sheet, int) else next((s for s in snapshot['sheets'] if s['name']==selection.sheet), None)
        if candidate is None:
            raise ValueError('候选工作表不存在')
        populated = [tuple(map(int,k.split(','))) for k,c in candidate['cells'].items() if c.get('value') is not None or c.get('formula')]
        last_row = max((r for r,c in populated), default=0); last_column = max((c for r,c in populated), default=0)
        if (last_row+1)*(last_column+1) > 2_000_000:
            raise ValueError('候选数据分布范围过大，请采用后选择实际数据区域处理')
        rows = []
        for r in range(last_row+1):
            row = []
            for c in range(last_column+1):
                value = candidate['cells'].get(f'{r},{c}', {})
                if value.get('formula') and value.get('result_state') != 'ready':
                    raise ValueError('候选包含尚未计算的公式，请采用后在编辑器中计算')
                row.append('' if value.get('value') is None else value['value'])
            rows.append(row)
        raw, sheet_name, info = pd.DataFrame(rows), candidate['name'], {**info,'validated':False}
    elif path.suffix == ".csv":
        try:
            raw = pd.read_csv(path, header=None, dtype=str, keep_default_na=False, encoding="utf-8-sig", skip_blank_lines=False)
        except UnicodeDecodeError:
            raw = pd.read_csv(path, header=None, dtype=str, keep_default_na=False, encoding="gb18030", skip_blank_lines=False)
        sheet_name = "CSV"
    elif path.suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        cached = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.worksheets[selection.sheet] if isinstance(selection.sheet, int) else book[selection.sheet]
            values = cached[sheet.title]
            sheet_name = sheet.title
            rows = []
            for formula_row, value_row in zip(sheet.iter_rows(), values.iter_rows()):
                row = []
                for cell, cached_cell in zip(formula_row, value_row):
                    value = cell.value
                    if cell.data_type == "f":
                        if cached_cell.value is None:
                            raise ValueError(f"{sheet_name}!{cell.coordinate} 公式没有缓存值，请先用 Excel 重算并保存")
                        value = cached_cell.value
                    if isinstance(value, (int, float)) and re.fullmatch(r"0{2,}", cell.number_format or ""):
                        if int(value) != value:
                            raise ValueError(f"{sheet_name}!{cell.coordinate} 使用编号格式但值不是整数，请先确认")
                        value = str(int(value)).zfill(len(cell.number_format))
                    row.append("" if value is None else value)
                rows.append(row)
            raw = pd.DataFrame(rows)
        finally:
            book.close()
            cached.close()
    else:
        raw = pd.read_excel(path, sheet_name=selection.sheet, header=None, dtype=object, keep_default_na=False, engine="calamine")
        sheet_name = str(selection.sheet)
    if header >= len(raw):
        raise ValueError("表头行超出数据范围")
    columns = [str(v).strip() for v in raw.iloc[header].tolist()]
    if not columns or any(not c or c == "nan" for c in columns) or len(set(columns)) != len(columns):
        raise ValueError("表头存在空白或重复列，请选择正确表头或整理源文件")
    if set(columns) & set(SOURCE) and not info.get("validated"):
        raise ValueError("输入包含系统保留的来源列")
    frame = raw.iloc[header + 1:].copy()
    frame.index = pd.RangeIndex(len(frame))
    frame.columns = columns
    if not info.get("validated"):
        frame[SOURCE[0]] = info["name"]
        frame[SOURCE[1]] = sheet_name
        frame[SOURCE[2]] = range(row_offset + selection.header_row + 1, row_offset + selection.header_row + 1 + len(frame))
    return frame


def preview(store, task, selection, offset=0, limit=50):
    if offset < 0 or not 1 <= limit <= 200:
        raise ValueError("分页范围无效，最多预览 200 行")
    frame = read(store, task, InputSelection.model_validate(selection))
    return {"columns": frame.columns.tolist(), "rows": records(frame.iloc[offset:offset + limit]), "total": len(frame)}


def require_columns(frame, columns):
    if not columns or not set(columns).issubset(frame.columns):
        raise ValueError("请选择存在的字段")


def aggregate_pivot(data, rows, columns, values, aggregate="sum"):
    """Aggregate ``data`` into a wide pivot frame for interactive preview.

    ``rows`` is a non-empty list of row fields, ``columns`` a single column field
    or None, ``values`` a non-empty list of numeric fields. Returns a flat,
    reset-indexed DataFrame whose columns are strings.
    """
    rows = list(rows or [])
    columns = columns or None
    values = list(values or [])
    if not rows:
        raise ValueError("请选择至少一个行字段")
    if not values:
        raise ValueError("请选择至少一个数值字段")
    require_columns(data, rows + ([columns] if columns else []) + values)
    if aggregate not in {"sum", "count", "min", "max", "mean"}:
        raise ValueError("请选择有效汇总规则")
    numeric = data.copy()
    for value in values:
        numeric[value] = pd.to_numeric(numeric[value], errors="raise")
    value_arg = values[0] if len(values) == 1 else values
    result = numeric.pivot_table(index=rows, columns=columns, values=value_arg, aggfunc=aggregate, fill_value=0).reset_index()
    if isinstance(result.columns, pd.MultiIndex):
        result.columns = [c if not isinstance(c, tuple) else " · ".join(str(x) for x in c if str(x) != "") for c in result.columns]
    else:
        result.columns = [str(c) for c in result.columns]
    return result


def run_operation(store, task_id, operation):
    started = time.perf_counter()
    op = Operation.model_validate(operation)
    if op.kind in {'create_table', 'edit_workbook'}:
        from .workbook_operations import run_workbook_operation
        return run_workbook_operation(store, task_id, op)
    if op.kind == "recalculate":
        from .excel_com import recalculate
        info = recalculate(store, task_id, op.inputs[0].file_id)
        from .workbook_results import register_file_result
        register_file_result(store, task_id, op, info)
        return info
    frames = [read(store, task_id, item) for item in (op.inputs[:1] if op.kind == "template" else op.inputs)]
    read_finished = time.perf_counter()
    p = op.params
    data = frames[0]
    issues = []
    extras = {}
    config = {"schema_version": 1}
    if op.kind in {'calculate', 'classify'}:
        from .rules import transform
        result, issues = transform(data, op.kind, p)
    elif op.kind == "append":
        result = merge_dataframes(frames, parse_config(config)).data
    elif op.kind == "join":
        if len(frames) != 2:
            raise ValueError("关联需要两个文件")
        keys = p.get("keys", [])
        for frame in frames:
            require_columns(frame, keys)
            if frame[keys].isna().any().any() or frame[keys].eq("").any().any():
                raise ValueError("关联键含空值，请先清洗，避免错误匹配")
        validation = p.get("validate", "many_to_one")
        if validation == "many_to_many":
            raise ValueError("首版禁止多对多直接关联，请先按业务规则汇总或去重")
        merged = merge_dataframes(frames, parse_config({**config, "merge": {"mode": "join", "keys": keys, "how": p.get("how", "left"), "validate": validation}}))
        result = merged.data
        extras = {"左表未匹配": merged.unmatched_left, "右表未匹配": merged.unmatched_right}
    elif op.kind == "clean":
        config.update(p.get("config", {}))
        config["input"] = {**config.get("input", {}), "header_row": op.inputs[0].header_row}
        if op.inputs[0].workbook_id and len(data):
            config['input']['header_row'] = int(data[SOURCE[2]].iloc[0]) - 1
        source_info = {'name': str(data[SOURCE[0]].iloc[0]) if len(data) else '工作簿'} if op.inputs[0].workbook_id else store.resolve_file(task_id, op.inputs[0].file_id)[1]
        cleaned = clean_dataframe(data, parse_config(config), source_info["name"])
        result = cleaned.data
        issues = [i.to_dict() for i in cleaned.issues]
    elif op.kind == "compare":
        if len(frames) != 2:
            raise ValueError("对账需要两个文件")
        keys = p.get("keys", [])
        fields = p.get("columns") or [c for c in data.columns if c not in SOURCE + keys and c in frames[1].columns]
        compared = compare_dataframes(*frames, parse_config({**config, "compare": {"keys": keys, "columns": fields}}))
        result = compared.changed
        extras = {"新增": compared.added, "删除": compared.removed, "左表来源": frames[0], "右表来源": frames[1]}
    elif op.kind == "group":
        keys, columns = p.get("keys", []), p.get("columns", [])
        require_columns(data, keys + columns)
        numeric = data.copy()
        for col in columns:
            numeric[col] = pd.to_numeric(numeric[col], errors="raise")
        aggregate = p.get("aggregate", "sum")
        if aggregate not in {"sum", "count", "min", "max", "mean"}:
            raise ValueError("不支持的汇总方式")
        result = numeric.groupby(keys, dropna=False, sort=False)[columns].agg(aggregate).reset_index()
        extras["来源明细"] = data
    elif op.kind == "melt":
        keys, columns = p.get("keys", []), p.get("columns", [])
        require_columns(data, keys + columns)
        result = data.melt(id_vars=list(dict.fromkeys(keys + [c for c in SOURCE if c in data.columns])), value_vars=columns, var_name=p.get("variable", "项目"), value_name=p.get("value", "数量"))
    elif op.kind == "pivot":
        values = p.get("values")
        if values is None:
            value = p.get("value")
            values = [value] if value else []
        rows = p.get("rows") or p.get("keys", [])
        column = p.get("columns") or p.get("column")
        if isinstance(column, list):
            column = column[0] if column else None
        result = aggregate_pivot(data, rows, column, values, p.get("aggregate", "sum"))
        extras["来源明细"] = data
    elif op.kind == "template":
        info = fill_template(store, task_id, op, frames[0])
        from .workbook_results import register_file_result
        register_file_result(store, task_id, op, info)
        return info
    elif op.kind == "report":
        return make_report(store, task_id, frames[0])
    elif op.kind == "chart":
        return run_chart(store, task_id, frames[0], op.params)
    elif op.kind == "pivot_table":
        return run_pivot_table(store, task_id, frames[0], op.params)
    elif op.kind == "sql":
        from data_toolkit.sql_runner import run_query, coerce_numeric_columns
        aliases = p.get("aliases", [])
        query = p.get("query", "")
        if not query or not isinstance(aliases, list) or len(aliases) != len(frames):
            raise ValueError("SQL 关联需要 query 与等量的表别名")
        for alias in aliases:
            if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", str(alias)):
                raise ValueError(f"表别名须为英文标识符: {alias}")
        result = run_query({alias: coerce_numeric_columns(frame) for alias, frame in zip(aliases, frames)}, query)
        stats = {"input_tables": list(aliases), "input_rows": [len(f) for f in frames], "output_rows": len(result)}
        return publish(store, task_id, result, {}, stats, [])
    else:
        raise ValueError("操作不支持")
    if issues:
        extras["问题明细"] = pd.DataFrame(issues)
    stats = {"input_rows": [len(f) for f in frames], "output_rows": len(result), "issues": len(issues)}
    stats['timings_ms'] = {'read': round((read_finished - started) * 1000), 'transform': round((time.perf_counter() - read_finished) * 1000)}
    info = publish(store, task_id, result, extras, stats, issues)
    from .workbook_results import register_result
    register_result(store, task_id, op, frames, result, info)
    return info


def publish(store, task_id, frame, extras, stats, issues):
    started = time.perf_counter()
    output_id = uuid.uuid4().hex
    folder = store.directory(task_id) / "outputs"
    folder.mkdir(exist_ok=True)
    target = folder / f"{output_id}.xlsx"
    temp = folder / f".{output_id}.tmp.xlsx"
    try:
        with pd.ExcelWriter(temp, engine="xlsxwriter", engine_kwargs={"options": {"strings_to_formulas": False, "strings_to_urls": False}}) as writer:
            for title, table in {"结果": frame, **extras}.items():
                if len(table) > 1048575 or len(table.columns) > 16384:
                    raise ValueError("结果超过 Excel 单表容量，请拆分任务")
                table.to_excel(writer, sheet_name=title, index=False)
                ws = writer.sheets[title]
                ws.freeze_panes(1, 0)
                if len(table.columns):
                    ws.autofilter(0, 0, len(table), len(table.columns) - 1)
                    ws.set_column(0, len(table.columns) - 1, 20)
        check = load_workbook(temp, read_only=True)
        check.close()
        temp.rename(target)
    finally:
        temp.unlink(missing_ok=True)
    info = {"id": output_id, "name": "处理结果.xlsx", "relative": str(target.relative_to(store.directory(task_id))), "statistics": stats, "issues": issues[:200], "validated": True}
    stats.setdefault('timings_ms', {})['export'] = round((time.perf_counter() - started) * 1000)
    task = store.get(task_id)
    task["outputs"].append(info)
    store.save(task)
    return info


def fill_template(store, task_id, op, frame):
    if len(op.inputs) != 2:
        raise ValueError("模板填写需依次选择数据文件与模板文件")
    path, _ = store.resolve_file(task_id, op.inputs[1].file_id)
    if path.suffix != ".xlsx":
        raise ValueError("模板仅支持 xlsx")
    book = load_workbook(path)
    output = store.directory(task_id) / "outputs" / f"{uuid.uuid4().hex}.xlsx"
    output.parent.mkdir(exist_ok=True)
    try:
        sheet = book[op.params.get("sheet") or book.sheetnames[0]]
        start = int(op.params.get("start_row", 2))
        mapping = op.params.get("mapping", {})
        if not mapping or start < 1 or start + len(frame) > 1048577:
            raise ValueError("请指定列映射和有效起始行")
        from openpyxl.utils.cell import column_index_from_string
        for name, letter in mapping.items():
            require_columns(frame, [name])
            column = column_index_from_string(letter)
            if column > 16384:
                raise ValueError("目标列超出 Excel 范围")
            for row, value in enumerate(frame[name], start=start):
                cell = sheet.cell(row, column)
                if cell.data_type == "f":
                    raise ValueError(f"目标 {cell.coordinate} 已有公式，拒绝覆盖")
                cell.value = None if pd.isna(value) else value
                if isinstance(value, str):
                    cell.data_type = "s"
        book.save(output)
    finally:
        book.close()
    check = load_workbook(output, read_only=True)
    check.close()
    info = {"id": uuid.uuid4().hex, "name": "模板结果.xlsx", "relative": str(output.relative_to(store.directory(task_id))), "statistics": {"written_rows": len(frame)}, "validated": True, "issues": [{"消息": "结构检查通过；公式重算与复杂对象需 Excel 2016 实机验收"}]}
    task = store.get(task_id)
    task["outputs"].append(info)
    store.save(task)
    return info


def make_report(store, task_id, frame):
    """Generate a data-quality report xlsx (summary / field profile) without a model."""
    from data_toolkit.profiling import profile_dataframe
    from data_toolkit.reporting import write_report
    columns = [c for c in frame.columns if not str(c).startswith('__source_')]
    data = frame[columns].replace(r'^\s*$', None, regex=True)
    summary = {'行数': len(data), '字段数': len(columns), '空值单元格': int(data.isna().sum().sum()), '重复行': int(data.duplicated().sum())}
    output_id = uuid.uuid4().hex
    folder = store.directory(task_id) / 'outputs'
    folder.mkdir(exist_ok=True)
    target = folder / f'{output_id}.xlsx'
    write_report(target, summary=summary, field_profile=profile_dataframe(data), issues=None)
    info = {'id': output_id, 'name': '数据质量报告.xlsx', 'relative': str(target.relative_to(store.directory(task_id))), 'statistics': summary, 'validated': True, 'issues': []}
    task = store.get(task_id)
    task['outputs'].append(info)
    store.save(task)
    return info


def run_chart(store, task_id, frame, spec):
    """Aggregate a frame into a chart: native Excel chart plus inline data for preview."""
    chart_type = spec.get('chart_type', 'column')
    if chart_type not in CHART_TYPES:
        raise ValueError('不支持的图表类型')
    aggregate = spec.get('aggregate', 'sum')
    if aggregate not in CHART_AGGREGATES:
        raise ValueError('不支持的聚合方式')
    category = spec.get('category')
    values = spec.get('values') or []
    if not category or not values:
        raise ValueError('请指定分类字段与至少一个数值字段')
    if chart_type == 'pie' and len(values) != 1:
        raise ValueError('饼图需要一个数值字段')
    columns = [c for c in frame.columns if not str(c).startswith('__source_')]
    data = frame[columns]
    require_columns(data, [category] + list(values))
    numeric = data.copy()
    for value in values:
        numeric[value] = pd.to_numeric(numeric[value], errors='raise')
    result = numeric.groupby(category, dropna=False, sort=False)[values].agg(aggregate).reset_index()
    truncated = len(result) > 200
    if truncated:
        result = result.sort_values(values[0], ascending=False).head(200)
    chart = {'type': chart_type, 'category': category, 'labels': result[category].tolist(), 'series': [{'name': str(v), 'values': result[v].tolist()} for v in values]}
    title = spec.get('title')
    if title:
        chart['title'] = title
    output_id = uuid.uuid4().hex
    folder = store.directory(task_id) / 'outputs'
    folder.mkdir(exist_ok=True)
    target = folder / f'{output_id}.xlsx'
    with pd.ExcelWriter(target, engine='xlsxwriter', engine_kwargs={'options': {'strings_to_formulas': False, 'strings_to_urls': False}}) as writer:
        result.to_excel(writer, index=False, sheet_name='图表数据')
        native = writer.book.add_chart({'type': chart_type})
        last = len(result)
        for index, value in enumerate(values):
            native.add_series({'name': str(value), 'categories': ['图表数据', 1, 0, last, 0], 'values': ['图表数据', 1, index + 1, last, index + 1]})
        if title:
            native.set_title({'name': title})
        writer.book.add_worksheet('图表').insert_chart('B2', native)
    check = load_workbook(target, read_only=True)
    check.close()
    info = {'id': output_id, 'name': '图表.xlsx', 'relative': str(target.relative_to(store.directory(task_id))), 'statistics': {'input_rows': len(data), 'output_rows': len(result), 'category': category, 'values': values, 'chart_type': chart_type, 'truncated': truncated}, 'chart': chart, 'validated': True, 'issues': []}
    task = store.get(task_id)
    task['outputs'].append(info)
    store.save(task)
    return info


def run_pivot_table(store, task_id, frame, spec):
    """Generate a native Excel pivot table (xl/pivotTables/ + cache parts)."""
    from .pivot_excel import build_pivot_xlsx
    output_id = uuid.uuid4().hex
    folder = store.directory(task_id) / 'outputs'
    folder.mkdir(exist_ok=True)
    target = folder / f'{output_id}.xlsx'
    build_pivot_xlsx(frame, spec, target)
    stats = {'input_rows': len(frame), 'rows': spec.get('rows'), 'columns': spec.get('columns'),
             'values': spec.get('values'), 'name': spec.get('name')}
    info = {'id': output_id, 'name': '透视表.xlsx', 'relative': str(target.relative_to(store.directory(task_id))),
            'statistics': stats, 'validated': True,
            'issues': [{'消息': '已生成原生透视表；请用 Excel 打开核对字段与汇总结果（结构检查通过，公式与对象需 Excel 实机验收）'}]}
    task = store.get(task_id)
    task['outputs'].append(info)
    store.save(task)
    return info
