"""Disposable, task-scoped preview indexes. Business execution never trusts this cache."""
import csv
from contextlib import closing
import hashlib
import json
import multiprocessing
import os
from pathlib import Path
import re
import sqlite3
import threading
import time

VERSION = 1
SOURCE = ['__source_file', '__source_sheet', '__source_row']


def stream_rows(path, info, selection):
    """Yield header then rows, preserving the business reader's Excel semantics."""
    header = selection['header_row']
    sheet_name = 'CSV'
    if path.suffix == '.xlsx':
        from openpyxl import load_workbook
        formulas = load_workbook(path, read_only=True, data_only=False)
        values = load_workbook(path, read_only=True, data_only=True)
        try:
            choice = selection['sheet']
            sheet = formulas.worksheets[choice] if isinstance(choice, int) else formulas[choice]
            sheet_name = sheet.title
            def excel_rows():
                for formula_row, value_row in zip(sheet.iter_rows(), values[sheet.title].iter_rows()):
                    row = []
                    for cell, cached in zip(formula_row, value_row):
                        value = cell.value
                        if cell.data_type == 'f':
                            if cached.value is None:
                                raise ValueError(f'{sheet_name}!{cell.coordinate} 公式没有缓存值，请先用 Excel 重算并保存')
                            value = cached.value
                        if isinstance(value, (int, float)) and re.fullmatch(r'0{2,}', cell.number_format or ''):
                            if int(value) != value:
                                raise ValueError(f'{sheet_name}!{cell.coordinate} 使用编号格式但值不是整数，请先确认')
                            value = str(int(value)).zfill(len(cell.number_format))
                        row.append('' if value is None else value)
                    yield row
            yield from normalize(excel_rows(), header, info, sheet_name)
        finally:
            formulas.close()
            values.close()
    elif path.suffix == '.csv':
        # Validate encoding in bounded chunks before selecting the fallback.
        encoding = 'utf-8-sig'
        try:
            with path.open(encoding=encoding) as handle:
                while handle.read(65536):
                    pass
        except UnicodeDecodeError:
            encoding = 'gb18030'
        with path.open(encoding=encoding, newline='') as handle:
            yield from normalize(csv.reader(handle), header, info, sheet_name)
    else:
        # Legacy .xls remains isolated; its existing reader owns its type rules.
        from .tables import read, records
        from .store import Store
        from .models import InputSelection
        frame = read(Store(Path(info['_root'])), info['_task'], InputSelection.model_validate(selection))
        yield frame.columns.tolist()
        for start in range(0, len(frame), 500):
            for record in records(frame.iloc[start:start + 500]):
                yield list(record.values())


def normalize(rows, header, info, sheet):
    columns = None
    for number, row in enumerate(rows, 1):
        if number < header:
            continue
        if number == header:
            columns = [str(value).strip() for value in row]
            if not columns or any(not c or c == 'nan' for c in columns) or len(set(columns)) != len(columns):
                raise ValueError('表头存在空白或重复列，请选择正确表头或整理源文件')
            if set(columns) & set(SOURCE) and not info.get('validated'):
                raise ValueError('输入包含系统保留的来源列')
            yield columns + ([] if info.get('validated') else SOURCE)
            continue
        if len(row) > len(columns):
            raise ValueError('数据列数超过表头列数，请确认表头')
        row = row + [''] * (len(columns) - len(row))
        yield row + ([] if info.get('validated') else [info['name'], sheet, number])
    if columns is None:
        raise ValueError('表头行超出数据范围')


def json_value(value):
    if hasattr(value, 'isoformat'):
        return value.isoformat()
    return str(value)


def index_worker(source, info, selection, temporary, final, error_path, budget):
    db = None
    rows = None
    try:
        rows = stream_rows(Path(source), info, selection)
        columns = next(rows)
        db = sqlite3.connect(temporary)
        db.execute('PRAGMA journal_mode=DELETE')
        db.execute('CREATE TABLE rows (id INTEGER PRIMARY KEY, body TEXT NOT NULL)')
        db.execute('CREATE TABLE metadata (columns TEXT, count INTEGER, complete INTEGER)')
        db.execute('INSERT INTO metadata VALUES (?,0,0)', (json.dumps(columns, ensure_ascii=False),))
        db.commit()
        batch, count = [], 0
        for row in rows:
            count += 1
            batch.append((count, json.dumps(row, ensure_ascii=False, default=json_value, allow_nan=False)))
            if len(batch) >= 200:
                write_batch(db, batch, count, temporary, budget)
                batch = []
        write_batch(db, batch, count, temporary, budget)
        db.execute('UPDATE metadata SET complete=1')
        db.commit()
        db.close()
        db = None
        for attempt in range(100):
            try:
                Path(temporary).replace(final)
                break
            except PermissionError:
                if attempt == 99:
                    raise
                time.sleep(.02)
    except Exception as exc:
        if db:
            db.close()
        error_temp = Path(str(error_path) + '.tmp')
        error_temp.write_text(json.dumps({'state': 'failed', 'message': str(exc)}, ensure_ascii=False), encoding='utf-8')
        error_temp.replace(error_path)
        Path(temporary).unlink(missing_ok=True)
    finally:
        if rows is not None:
            rows.close()


def write_batch(db, batch, count, path, budget):
    db.executemany('INSERT INTO rows VALUES (?,?)', batch)
    db.execute('UPDATE metadata SET count=?', (count,))
    db.commit()
    if Path(path).stat().st_size > budget:
        raise ValueError('预览缓存容量已达上限，请清理历史缓存或直接执行处理')


class PreviewService:
    def __init__(self, store, quota=1024 ** 3):
        self.store, self.quota = store, quota
        self.lock = threading.RLock()
        self.process = None
        self.active = None
        self.requests = {}
        self.closed = False
        for path in self.store.root.glob('tasks/*/preview-cache/*.partial*'):
            path.unlink(missing_ok=True)

    def start(self, task_id, selection):
        from .models import InputSelection
        choice = InputSelection.model_validate(selection).model_dump()
        source, info = self.store.resolve_file(task_id, choice['file_id'])
        signature = [source.stat().st_size, source.stat().st_mtime_ns]
        key = hashlib.sha256(json.dumps([VERSION, task_id, choice, signature], sort_keys=True).encode()).hexdigest()
        with self.lock:
            if self.closed:
                raise ValueError('预览服务已关闭')
            # A worker may still be exiting after publishing its completed cache.
            # Revalidate that file on restart, including if it was corrupted meanwhile.
            if self.active == key and self.process and self.process.is_alive() and key in self.requests and not self.requests[key][0].exists():
                return {'preview_id': key}
            self.cancel()
            folder = self.store.directory(task_id) / 'preview-cache'
            folder.mkdir(exist_ok=True)
            final = folder / f'{key}.sqlite'
            temporary = folder / f'{key}.partial'
            error = folder / f'{key}.error.json'
            self.requests[key] = (final, temporary, error, source, signature)
            # IDs are disposable; only the newest 100 view requests stay addressable.
            while len(self.requests) > 100:
                self.requests.pop(next(iter(self.requests)))
            if final.exists():
                try:
                    with closing(sqlite3.connect(f'{final.as_uri()}?mode=ro', uri=True)) as db:
                        if db.execute('PRAGMA quick_check').fetchone()[0] != 'ok':
                            raise sqlite3.DatabaseError()
                        if db.execute('SELECT complete FROM metadata').fetchone()[0] != 1:
                            raise sqlite3.DatabaseError()
                    os.utime(final, None)
                    return {'preview_id': key}
                except sqlite3.DatabaseError:
                    final.unlink(missing_ok=True)
            error.unlink(missing_ok=True)
            caches = sorted(self.store.root.glob('tasks/*/preview-cache/*.sqlite'), key=lambda p: p.stat().st_mtime)
            used = sum(p.stat().st_size for p in caches)
            # Keep room for this index. Evict only disposable completed caches.
            reserve = min(self.quota, max(source.stat().st_size * 8, 16 * 1024 ** 2))
            while caches and self.quota - used < reserve:
                old = caches.pop(0)
                used -= old.stat().st_size
                old.unlink()
            self.process = multiprocessing.get_context('spawn').Process(target=index_worker, args=(
                str(source), {**info, '_root': str(self.store.root), '_task': task_id}, choice,
                str(temporary), str(final), str(error), self.quota - used))
            self.process.start()
            self.active = key
            return {'preview_id': key}

    def page(self, preview_id, offset=0, limit=50):
        if offset < 0 or not 1 <= limit <= 200:
            raise ValueError('分页范围无效，最多预览 200 行')
        with self.lock:
            if preview_id not in self.requests:
                raise ValueError('预览已过期，请重新选择文件')
            final, temporary, error, source, signature = self.requests[preview_id]
            if not source.exists() or [source.stat().st_size, source.stat().st_mtime_ns] != signature:
                return {'state': 'stale', 'message': '文件已变化，请重新加载预览'}
            if error.exists():
                return json.loads(error.read_text(encoding='utf-8'))
            for path in (final, temporary, final):
                if not path.exists():
                    continue
                try:
                    with closing(sqlite3.connect(f'{path.as_uri()}?mode=ro', uri=True, timeout=0.1)) as db:
                        db.execute('BEGIN')
                        metadata = db.execute('SELECT * FROM metadata').fetchone()
                        if metadata is None and path == temporary:
                            # Table creation is committed before the first metadata row.
                            # A reader during this interval must keep waiting for the worker.
                            continue
                        columns, count, complete = metadata
                        columns = json.loads(columns)
                        rows = [dict(zip(columns, json.loads(body))) for (body,) in db.execute(
                            'SELECT body FROM rows WHERE id>? ORDER BY id LIMIT ?', (offset, limit))]
                    complete = complete and path == final
                    if complete:
                        os.utime(path, None)
                    return {'state': 'ready' if complete else 'indexing', 'columns': columns, 'rows': rows,
                            'indexed_rows': count, 'total': count if complete else None}
                except sqlite3.OperationalError:
                    continue
                except (sqlite3.DatabaseError, ValueError, TypeError):
                    return {'state': 'failed', 'message': '预览缓存损坏，请重新加载'}
            alive = self.active == preview_id and self.process and self.process.is_alive()
            return {'state': 'indexing' if alive else 'cancelled', 'columns': [], 'rows': [], 'indexed_rows': 0, 'total': None}

    def cancel(self, preview_id=None):
        with self.lock:
            if preview_id and preview_id != self.active:
                return False
            if self.process:
                if self.process.is_alive():
                    self.process.terminate()
                self.process.join(5)
                if self.process.is_alive():
                    self.process.kill()
                    self.process.join()
                self.process.close()
                self.process = None
            if self.active in self.requests:
                temporary = self.requests[self.active][1]
                temporary.unlink(missing_ok=True)
                Path(str(temporary) + '-journal').unlink(missing_ok=True)
            self.active = None
            return True

    def close(self):
        with self.lock:
            self.closed = True
            self.cancel()
