"""Task-scoped MCP server; stdout is reserved for the MCP transport."""
import argparse
import re
import sys
import anyio
from pathlib import Path
from mcp.server.fastmcp import FastMCP
from mcp.server.stdio import stdio_server
from .store import Store
from .models import InputSelection, Operation
from .tables import overview, preview, read, run_operation, publish, make_report, run_chart, run_pivot_table, emit_step
from .workbook_model import formula_issue
from data_toolkit.profiling import profile_columns
from data_toolkit.sql_runner import run_query, coerce_numeric_columns


class TaskMCP(FastMCP):
    async def run_stdio_async(self):
        # Keep the process-owned TextIOWrappers alive. Re-wrapping their buffers
        # lets GC close stdout before the frozen bootloader performs its flush.
        sys.stdin.reconfigure(encoding="utf-8", errors="replace")
        sys.stdout.reconfigure(encoding="utf-8")
        async with stdio_server(anyio.wrap_file(sys.stdin), anyio.wrap_file(sys.stdout)) as (reader, writer):
            await self._mcp_server.run(reader, writer, self._mcp_server.create_initialization_options())


def create_server(root: Path, task: str):
    store = Store(root)
    store.get(task)
    server = TaskMCP("datacraft_mcp")

    @server.tool()
    def datacraft_files() -> dict:
        """列出本任务导入文件及已生成文件的标识。仅能使用这些标识。"""
        record = store.get(task)
        from .workbooks import Workbooks
        service = Workbooks(store)
        books = service.list(task)
        for book in books:
            book['sheets'] = [{'id': s['id'], 'name': s['name']} for s in service.load(task, book['id'])['snapshot']['sheets']]
        return {"inputs": record["files"], "outputs": record["outputs"], 'workbooks': books}

    @server.tool()
    def datacraft_inspect(file_id: str) -> dict:
        """读取工作表目录；表头由用户确认，不猜测。"""
        path, info = store.resolve_file(task, file_id)
        if info.get('kind') == 'workbook_candidate':
            import json
            snapshot = json.loads(path.read_text(encoding='utf-8'))
            return {'sheets':[s['name'] for s in snapshot['sheets']],'note':'待用户整份核对的工作簿候选；模型不能采用或导出'}
        return overview(path)

    @server.tool()
    def datacraft_preview(selection: InputSelection, offset: int = 0, limit: int = 50) -> dict:
        """分页读取数据（最多 200 行），返回真实总行数与来源列。"""
        return preview(store, task, selection.model_dump(), offset, limit)

    @server.tool()
    def datacraft_execute(operation: Operation) -> dict:
        """生成候选结果，不代替用户采用或导出。join: keys/how/validate；group: keys/columns/aggregate；melt: keys/columns；pivot: keys/column/value/aggregate；clean: config(schema_version=1)；template: sheet/start_row/mapping。
        calculate params={columns:[{name:'金额',expression:{op:'multiply',args:[{field:'数量'},{field:'单价'}]}}]}，仅 add/subtract/multiply/divide/round/min/max/abs；round 第二参数为 {number:2}。
        classify params={columns:[{name:'等级',rules:[{when:{field:'金额',operator:'ge',value:100},label:'高'}],default:'普通'}]}。条件允许 all/any 数组及 eq/ne/gt/ge/lt/le。
        create_table inputs=[]，params={name,sheet_name,columns:['字段'],rows:[[null]],blank_rows:20}。只填用户明确提供的数据，缺失数字 null；用户明确要求示例才使用 example_data:true 并注明。
        edit_workbook 需要工作簿输入（workbook_id/version/sheet_id/range），params={edits:[{kind:'set_values',range:{r0:1,c0:0,r1:1,c1:0},values:[[1]]}]}。范围零基含结束。还允许 set_style(style:font_name/font_size/bold/italic/font_color/bg_color/align/valign/text_wrap/num_format/borders)、add_sheet(name)、rename_sheet(name)。不可执行任意内核命令。所有业务歧义先询问；不得扩大用户选区。"""
        result = run_operation(store, task, operation.model_dump())
        record = store.get(task)
        plan = record.get("plan") or {"steps": [], "questions": []}
        plan["steps"].append(operation.model_dump())
        record["plan"] = plan
        store.save(record)
        store.event(task, "output", result)
        emit_step(store, task, operation.kind, result)
        return result

    @server.tool()
    def datacraft_validate(file_id: str) -> dict:
        """验证本任务生成的 xlsx 可重新读取；不代替业务数值核对或 Excel 实机验收。"""
        from openpyxl import load_workbook
        path, info = store.resolve_file(task, file_id)
        if not info.get("validated"):
            raise ValueError("此文件不是已登记的输出")
        if info.get('kind') == 'workbook_candidate':
            from .workbooks import Workbooks
            return {'readable': True, 'candidate': True, 'review': Workbooks(store).review(task, file_id)}
        book = load_workbook(path, read_only=True)
        try:
            return {"readable": True, "sheets": book.sheetnames, "statistics": info.get("statistics", {})}
        finally:
            book.close()

    @server.tool()
    def datacraft_profile(selection: InputSelection, sample_limit: int = 5) -> dict:
        """读取字段画像：类型、空值率、唯一数、高频取值示例与数值范围。处理前先看画像，不要凭预览猜类型。"""
        if not 1 <= sample_limit <= 20:
            raise ValueError('取值示例数量须在 1–20 之间')
        frame = read(store, task, selection)
        columns = [c for c in frame.columns if not str(c).startswith('__source_')]
        return {"row_count": len(frame), "column_count": len(columns), "columns": profile_columns(frame[columns], sample_limit)}

    @server.tool()
    def datacraft_sql(query: str, tables: dict[str, InputSelection]) -> dict:
        """用只读 SQL 关联多张表。tables 形如 {'orders': 文件/工作簿选区}，query 引用这些别名；别名须为英文标识符。仅允许 SELECT/WITH 单查询。"""
        selections = {}
        for alias, selection in tables.items():
            if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', alias):
                raise ValueError(f'表别名须为英文标识符: {alias}')
            selections[alias] = InputSelection.model_validate(selection)
        frames = {alias: coerce_numeric_columns(read(store, task, selection)) for alias, selection in selections.items()}
        result = run_query(frames, query)
        stats = {'input_tables': list(frames), 'input_rows': [len(f) for f in frames.values()], 'output_rows': len(result)}
        info = publish(store, task, result, {}, stats, [])
        record = store.get(task)
        plan = record.get('plan') or {'steps': [], 'questions': []}
        plan['steps'].append({'kind': 'sql', 'params': {'query': query, 'tables': {a: s.model_dump(exclude_none=True) for a, s in selections.items()}}})
        record['plan'] = plan
        store.save(record)
        store.event(task, 'output', info)
        emit_step(store, task, 'sql', info)
        return info

    @server.tool()
    def datacraft_report(selection: InputSelection) -> dict:
        """生成数据质量报告 xlsx（汇总/字段概况）；不替代业务数值核对。"""
        frame = read(store, task, selection)
        info = make_report(store, task, frame)
        store.event(task, 'output', info)
        emit_step(store, task, 'report', info)
        return info

    @server.tool()
    def datacraft_chart(selection: InputSelection, chart_type: str, category: str, values: list[str], aggregate: str = 'sum', title: str = '') -> dict:
        """按 category 分组对 values 聚合，生成图表（前端预览 + 原生 Excel 图表）。chart_type: column/bar/line/area/pie；aggregate: sum/count/min/max/mean；pie 只能一个 values。数值字段须可转数字。"""
        frame = read(store, task, selection)
        info = run_chart(store, task, frame, {'chart_type': chart_type, 'category': category, 'values': values, 'aggregate': aggregate, 'title': title})
        store.event(task, 'output', info)
        emit_step(store, task, 'chart', info)
        return info

    @server.tool()
    def datacraft_pivot_table(selection: InputSelection, rows: list[str], columns: list[str], values: list[dict], name: str = '') -> dict:
        """生成原生 Excel 透视表（可被 Excel 打开并继续交互）。rows 至少 1 个行字段；columns 最多 1 个列字段；values 恰好 1 个 {field, aggregate}，aggregate: sum/count/min/max/mean；name 为透视表名称。数值字段须可转数字。"""
        frame = read(store, task, selection)
        spec = {'rows': rows, 'columns': columns, 'values': values, 'name': name}
        info = run_pivot_table(store, task, frame, spec)
        store.event(task, 'output', info)
        emit_step(store, task, 'pivot_table', info)
        return info

    @server.tool()
    def datacraft_audit() -> dict:
        """汇总本任务的处理步骤与产出，说明数据来源（血缘）；只读。"""
        record = store.get(task)
        return {
            'steps': (record.get('plan') or {}).get('steps', []),
            'inputs': record.get('files', []),
            'outputs': [{'id': o.get('id'), 'name': o.get('name'), 'kind': o.get('kind', 'xlsx'), 'statistics': o.get('statistics', {})} for o in record.get('outputs', [])],
        }

    @server.tool()
    def datacraft_formula_generate(formulas: list[str]) -> dict:
        """校验候选公式（白名单函数、无数组/外部引用、长度限制）；不写入。返回逐条 ok/issue 供修正。"""
        checked = []
        for formula in formulas:
            issue = formula_issue(formula)
            checked.append({'formula': formula, 'ok': issue is None, 'issue': issue})
        return {'candidates': checked}

    @server.tool()
    def datacraft_formula_explain(selection: InputSelection) -> dict:
        """读取指定区域单元格的公式并给出地址与当前值，供逐格解释；仅读不写。"""
        if not selection.workbook_id:
            raise ValueError('公式解释需要工作簿输入')
        if not selection.range:
            raise ValueError('请选择要解释的公式单元格区域')
        from .workbooks import Workbooks
        from openpyxl.utils import get_column_letter
        record = Workbooks(store).load(task, selection.workbook_id, selection.version)
        sheet = next((s for s in record['snapshot']['sheets'] if s['id'] == selection.sheet_id), None)
        if sheet is None:
            raise ValueError('工作表不存在')
        area = selection.range
        cells = []
        for r in range(area['r0'], area['r1'] + 1):
            for c in range(area['c0'], area['c1'] + 1):
                cell = sheet['cells'].get(f'{r},{c}')
                if cell and cell.get('formula'):
                    cells.append({'address': f"{sheet['name']}!{get_column_letter(c + 1)}{r + 1}", 'formula': cell['formula'], 'value': cell.get('value')})
        return {'cells': cells}

    return server


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", required=True)
    parser.add_argument("--task", required=True)
    args = parser.parse_args()
    create_server(Path(args.root), args.task).run(transport="stdio")


if __name__ == "__main__":
    main()
