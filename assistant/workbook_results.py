"""Bridge deterministic data results into workbook candidates and complete paged review."""
import copy
import json
import pandas as pd

from .workbooks import Workbooks, workbook_changes
from .workbook_excel import import_workbook
from .workbook_model import MAX_CELLS, validate_snapshot, get_sheet, invalidate_formulas
from .workbook_operations import cell


def file_rows(path):
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        yield {'kind': 'structure', 'location': '结果工作表', 'old': None, 'new': book.sheetnames, 'note': '完整文件结果；本次变化以原输入文件和此副本为依据。'}
        for sheet in book:
            for r, values in enumerate(sheet.iter_rows(values_only=True), 1):
                yield {'kind': 'result', 'location': f'{sheet.title}!第 {r} 行', 'old': None, 'new': [scalar(v) for v in values]}
    finally:
        book.close()


def register_file_result(store, task, op, info):
    from .workbook_model import formula_errors
    path, _ = store.resolve_file(task, info['id'])
    snapshot, limitations, before, needs_calculation = None, [], None, False
    try:
        snapshot, limitations = import_workbook(path)
        needs_calculation = bool(formula_errors(snapshot))
        original, _ = store.resolve_file(task, op.inputs[-1].file_id)
        before, _ = import_workbook(original)
        names = {s['name']:s['id'] for s in before['sheets']}
        for s in snapshot['sheets']:
            if s['name'] in names:s['id']=names[s['name']]
    except ValueError as exc:
        if '200000' not in str(exc):raise
        limitations.append(str(exc))
        snapshot = None
    changes = workbook_changes(before,snapshot) if snapshot is not None else file_rows(path)
    Workbooks(store).candidate(task, info['id'], None if limitations else snapshot, summary={'operation':op.kind,'inputs':[i.model_dump(exclude_none=True) for i in op.inputs],'rules':op.params,'limitations':limitations,'requires_excel_recalculation':needs_calculation and bool(limitations)}, changes=changes)


def scalar(value):
    if value is None or pd.isna(value):
        return None
    if hasattr(value, 'item'):
        value = value.item()
    return value if isinstance(value, (str, bool, int, float)) else str(value)


def data_changes(op, before, after):
    from .tables import SOURCE
    # Only row-preserving operations have a justified old -> new correspondence.
    aligned = op.kind in {'calculate', 'classify'} or op.kind == 'clean' and len(before) == len(after) and all(before[c].equals(after[c]) for c in SOURCE)
    if not aligned:
        yield {'kind': 'structure', 'location': '结果结构', 'old': {'rows': len(before), 'columns': list(before.columns)}, 'new': {'rows': len(after), 'columns': list(after.columns)}, 'note': '此操作改变记录结构；下列完整结果不表示与输入逐行对应。来源见操作输入及来源明细。'}
        for i, values in enumerate(after.itertuples(index=False, name=None), 2):
            yield {'kind': 'result', 'location': f'结果!第 {i} 行', 'old': None, 'new': {k: scalar(v) for k, v in zip(after.columns, values)}}
        return
    old_columns = list(before.columns)
    old_rows = before.itertuples(index=False, name=None)
    for previous, values in zip(old_rows, after.itertuples(index=False, name=None)):
        old, new = dict(zip(old_columns, previous)), dict(zip(after.columns, values))
        for column in after.columns:
            a, b = scalar(old.get(column)), scalar(new[column])
            if a != b:
                yield {'kind': 'value', 'location': f"{new.get(SOURCE[0], '')} / {new.get(SOURCE[1], '')}!第 {new.get(SOURCE[2], '')} 行 · {column}", 'old': a, 'new': b, 'source': {k: scalar(new.get(k)) for k in SOURCE}}


def register_result(store, task, op, frames, result, info):
    from .tables import SOURCE
    service, base, snapshot, changes = Workbooks(store), None, None, None
    selection = op.inputs[0]
    safe_count = sum((len(f) + 1)*len(f.columns) for f in [result])
    if safe_count <= MAX_CELLS:
        path, _ = store.resolve_file(task, info['id'])
        try:
            imported, limitations = import_workbook(path)
            if not limitations:
                snapshot = imported
        except ValueError as exc:
            if '200000' not in str(exc):
                raise
    if selection.workbook_id:
        base = {'id': selection.workbook_id, 'revision': selection.version}
        original = service.load(task, selection.workbook_id, selection.version)['snapshot']
        if snapshot and op.kind in {'calculate', 'classify', 'clean'} and len(result) == len(frames[0]):
            snapshot = copy.deepcopy(original)
            target = get_sheet(snapshot, selection.sheet_id)
            _, source = service.input(task, selection)
            area = source['area']
            data_fields = [c for c in result.columns if c not in SOURCE]
            old_fields = [c for c in frames[0].columns if c not in SOURCE]
            for ci, name in enumerate(data_fields, area['c0']):
                if name not in old_fields and any(k in target['cells'] and (target['cells'][k].get('value') is not None or target['cells'][k].get('formula')) for k in [f'{r},{ci}' for r in range(area['r0'], area['r1'] + 1)]):
                    raise ValueError('新增计算字段与选区右侧已有内容冲突，请先留出空列')
                header = area['r0'] + selection.header_row - 1
                if name not in old_fields:
                    target['cells'][f'{header},{ci}'] = cell(name)
                for index, value in enumerate(result[name], header + 1):
                    key = f'{index},{ci}'
                    previous = target['cells'].get(key, {})
                    # Unchanged cells retain formulas and styles.
                    if name in old_fields and scalar(value) == scalar(frames[0][name].iloc[index-header-1]):
                        continue
                    target['cells'][key] = {**cell(scalar(value)), **({'style': previous['style']} if previous.get('style') else {})}
            target['column_count'] = max(target['column_count'], area['c0'] + len(data_fields))
            invalidate_formulas(snapshot)
        elif snapshot:
            candidate = copy.deepcopy(original)
            used = {s['name'].casefold() for s in candidate['sheets']}
            for sheet in snapshot['sheets']:
                name, index = sheet['name'], 1
                while sheet['name'].casefold() in used:
                    index += 1; sheet['name'] = f'{name[:24]} ({index})'
                used.add(sheet['name'].casefold()); candidate['sheets'].append(sheet)
            snapshot = candidate
        if snapshot:
            try:
                validate_snapshot(snapshot)
            except ValueError as exc:
                if '200000' not in str(exc):
                    raise
                snapshot = None
            if snapshot:
                changes = workbook_changes(original, snapshot)
    if changes is None:
        changes = data_changes(op, frames[0], result)
    service.candidate(task, info['id'], snapshot, base, {'operation': op.kind, 'inputs': [i.model_dump(exclude_none=True) for i in op.inputs], 'rules': op.params, 'output_rows': len(result)}, changes)
