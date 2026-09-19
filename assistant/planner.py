"""Native function-calling agent loop: planner produces an OperationPlan, executor runs it deterministically.

This is the direct OpenAI-compatible backend (方向 5). Unlike ``agent.AgentRunner``, it
does not spawn the Qwen Code CLI or an MCP subprocess: it talks to the endpoint via
``model_registry.chat()`` and executes read-only observation tools in-process. The model
only emits text and structured ``Operation`` tool calls; execution stays in the
deterministic ``tables.run_operation`` kernel.
"""
import json
from typing import get_args

from .agent import ask_question
from .model_registry import chat, api_key as model_api_key
from .models import Operation, OperationPlan
from .tables import overview, preview, read, run_operation, emit_step

MAX_TURNS = 12

SYSTEM_PROMPT = (
    "你是内网 Excel 助手，只使用 datacraft 工具处理用户任务。"
    "先用 datacraft_files / datacraft_inspect / datacraft_preview / datacraft_profile 查看真实字段，"
    "再通过 datacraft_execute 逐步登记操作（operation），绝不猜测字段名。"
    "业务歧义调用 ask_user_question。生成公式先用 datacraft_formula_generate 校验。"
    "原生透视表用 datacraft_execute 的 kind='pivot_table'：params={rows:至少1个行字段, columns:最多1个列字段, values:[{field,aggregate}]恰好1个, name}，aggregate∈sum/count/min/max/mean，数值字段须可转数字。"
    "结果必须经过工具登记，不要编造结果文件；不得执行 Shell、安装依赖或编写脚本。"
    "用中文说明行数、异常和校验限制。"
)

OPERATION_KINDS = list(get_args(Operation.model_fields["kind"].annotation))


def _selection(description):
    """Flattened JSON schema for InputSelection (no $ref, tolerated by Ollama/vLLM)."""
    return {
        "type": "object",
        "description": description,
        "properties": {
            "file_id": {"type": "string", "description": "输入文件标识（来自 datacraft_files）"},
            "sheet": {"type": ["string", "integer"], "description": "工作表名或索引，默认 0"},
            "header_row": {"type": "integer", "description": "表头行号，默认 1"},
            "workbook_id": {"type": "string", "description": "工作簿标识（编辑/公式场景）"},
            "version": {"type": "integer", "description": "工作簿版本"},
            "sheet_id": {"type": "string", "description": "工作表标识"},
            "range": {"type": "object", "description": "选区 {r0,c0,r1,c1}，零基含结束"},
        },
    }


def _operation_schema():
    return {
        "type": "object",
        "description": "单个确定性数据操作。kind 决定语义；inputs 为输入；params 为具体参数。",
        "required": ["kind", "inputs", "params"],
        "properties": {
            "kind": {"type": "string", "enum": OPERATION_KINDS},
            "inputs": {"type": "array", "items": _selection("操作输入")},
            "params": {"type": "object"},
        },
    }


def _question_schema():
    return {
        "type": "object",
        "required": ["questions"],
        "properties": {
            "questions": {
                "type": "array",
                "items": {
                    "type": "object",
                    "required": ["question"],
                    "properties": {
                        "question": {"type": "string"},
                        "options": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "required": ["label", "value"],
                                "properties": {"label": {"type": "string"}, "value": {"type": "string"}},
                            },
                        },
                    },
                },
            },
        },
    }


def _tool(name, description, parameters):
    return {"type": "function", "function": {"name": name, "description": description, "parameters": parameters}}


class NativeAgent:
    def __init__(self, store, task, config):
        self.store, self.task, self.config = store, task, config
        self._healed = False
        self._final = ""

    # -- tool surface -------------------------------------------------------
    def tool_schemas(self):
        return [
            _tool("datacraft_files", "列出本任务导入文件及已生成文件的标识。仅能使用这些标识。", {"type": "object", "properties": {}}),
            _tool("datacraft_inspect", "读取工作表目录；表头由用户确认，不猜测。", {"type": "object", "required": ["file_id"], "properties": {"file_id": {"type": "string"}}}),
            _tool("datacraft_preview", "分页读取数据（最多 200 行），返回真实总行数与来源列。", {"type": "object", "required": ["selection"], "properties": {"selection": _selection("预览输入"), "offset": {"type": "integer"}, "limit": {"type": "integer"}}}),
            _tool("datacraft_profile", "读取字段画像：类型、空值率、唯一数、高频取值示例与数值范围。处理前先看画像。", {"type": "object", "required": ["selection"], "properties": {"selection": _selection("画像输入"), "sample_limit": {"type": "integer"}}}),
            _tool("datacraft_formula_generate", "校验候选公式（白名单函数、无数组/外部引用、长度限制）；不写入。", {"type": "object", "required": ["formulas"], "properties": {"formulas": {"type": "array", "items": {"type": "string"}}}}),
            _tool("datacraft_formula_explain", "读取指定区域单元格的公式并给出地址与当前值；仅读不写。", {"type": "object", "required": ["selection"], "properties": {"selection": _selection("公式区域输入")}}),
            _tool("datacraft_audit", "汇总本任务的处理步骤与产出，说明数据来源（血缘）；只读。", {"type": "object", "properties": {}}),
            _tool("ask_user_question", "遇到业务歧义时向用户提问，逐条列出问题与可选答案。", _question_schema()),
            _tool("datacraft_execute", "登记一个确定性操作步骤（不立即执行，仅列入计划）。生成候选结果，不代替用户采用或导出。", {"type": "object", "required": ["operation"], "properties": {"operation": _operation_schema()}}),
        ]

    # -- observation handlers ----------------------------------------------
    def _observe(self, name, args):
        if name == "datacraft_files":
            return self._files()
        if name == "datacraft_inspect":
            return self._inspect(args.get("file_id"))
        if name == "datacraft_preview":
            return preview(self.store, self.task, args.get("selection") or {}, args.get("offset", 0), args.get("limit", 50))
        if name == "datacraft_profile":
            return self._profile(args.get("selection") or {}, args.get("sample_limit", 5))
        if name == "datacraft_formula_generate":
            return self._formula_generate(args.get("formulas") or [])
        if name == "datacraft_formula_explain":
            return self._formula_explain(args.get("selection") or {})
        if name == "datacraft_audit":
            return self._audit()
        raise ValueError(f"未知的观测工具: {name}")

    def _files(self):
        from .workbooks import Workbooks
        record = self.store.get(self.task)
        service = Workbooks(self.store)
        books = service.list(self.task)
        for book in books:
            book["sheets"] = [{"id": s["id"], "name": s["name"]} for s in service.load(self.task, book["id"])["snapshot"]["sheets"]]
        return {"inputs": record["files"], "outputs": record["outputs"], "workbooks": books}

    def _inspect(self, file_id):
        path, info = self.store.resolve_file(self.task, file_id)
        if info.get("kind") == "workbook_candidate":
            snapshot = json.loads(path.read_text(encoding="utf-8"))
            return {"sheets": [s["name"] for s in snapshot["sheets"]], "note": "待用户整份核对的工作簿候选；模型不能采用或导出"}
        return overview(path)

    def _profile(self, selection, sample_limit=5):
        from data_toolkit.profiling import profile_columns
        from .models import InputSelection
        if not 1 <= sample_limit <= 20:
            raise ValueError("取值示例数量须在 1–20 之间")
        frame = read(self.store, self.task, InputSelection.model_validate(selection))
        columns = [c for c in frame.columns if not str(c).startswith("__source_")]
        return {"row_count": len(frame), "column_count": len(columns), "columns": profile_columns(frame[columns], sample_limit)}

    def _formula_generate(self, formulas):
        from .workbook_model import formula_issue
        checked = []
        for formula in formulas:
            issue = formula_issue(formula)
            checked.append({"formula": formula, "ok": issue is None, "issue": issue})
        return {"candidates": checked}

    def _formula_explain(self, selection):
        from .workbooks import Workbooks
        from openpyxl.utils import get_column_letter
        from .models import InputSelection
        selection = InputSelection.model_validate(selection)
        if not selection.workbook_id:
            raise ValueError("公式解释需要工作簿输入")
        if not selection.range:
            raise ValueError("请选择要解释的公式单元格区域")
        record = Workbooks(self.store).load(self.task, selection.workbook_id, selection.version)
        sheet = next((s for s in record["snapshot"]["sheets"] if s["id"] == selection.sheet_id), None)
        if sheet is None:
            raise ValueError("工作表不存在")
        area = selection.range
        cells = []
        for r in range(area["r0"], area["r1"] + 1):
            for c in range(area["c0"], area["c1"] + 1):
                cell = sheet["cells"].get(f"{r},{c}")
                if cell and cell.get("formula"):
                    cells.append({"address": f"{sheet['name']}!{get_column_letter(c + 1)}{r + 1}", "formula": cell["formula"], "value": cell.get("value")})
        return {"cells": cells}

    def _audit(self):
        record = self.store.get(self.task)
        return {
            "steps": (record.get("plan") or {}).get("steps", []),
            "inputs": record.get("files", []),
            "outputs": [{"id": o.get("id"), "name": o.get("name"), "kind": o.get("kind", "xlsx"), "statistics": o.get("statistics", {})} for o in record.get("outputs", [])],
        }

    # -- planning loop ------------------------------------------------------
    def _chat(self, messages, tools):
        if not (self.config.get("base_url") or "").strip() or not (self.config.get("model") or "").strip():
            raise ValueError("请先配置内网模型服务地址与 model 标识")
        return chat(self.config["base_url"], self.config["model"], messages,
                    api_key=model_api_key(self.config.get("api_key_env", "EXCEL_ASSISTANT_API_KEY")),
                    tools=tools)

    def _plan(self, messages):
        tools = self.tool_schemas()
        pending = []
        for _ in range(MAX_TURNS):
            response = self._chat(messages, tools)
            assistant = {"role": "assistant", "content": response["content"]}
            if response["tool_calls"]:
                assistant["tool_calls"] = [
                    {"id": c["id"], "type": "function",
                     "function": {"name": c["name"], "arguments": c["raw_arguments"] if isinstance(c["raw_arguments"], str) else json.dumps(c["raw_arguments"] or {})}}
                    for c in response["tool_calls"]
                ]
            messages.append(assistant)
            if response["content"]:
                self._final = response["content"]
            if not response["tool_calls"]:
                break
            for call in response["tool_calls"]:
                name, args = call["name"], call["arguments"] or {}
                if name == "datacraft_execute":
                    pending.append(args.get("operation") or {})
                    result_text = "已列入计划"
                elif name == "ask_user_question":
                    answered = ask_question(self.store, self.task, args)
                    result_text = json.dumps({"answers": answered.get("answers") if answered else None}, ensure_ascii=False)
                else:
                    try:
                        result_text = json.dumps(self._observe(name, args), ensure_ascii=False, default=str)
                    except Exception as exc:
                        result_text = json.dumps({"error": str(exc)}, ensure_ascii=False)
                messages.append({"role": "tool", "tool_call_id": call["id"], "content": result_text})
        if not pending:
            raise ValueError("模型尚未产出任何操作步骤，请补充要求后继续")
        return OperationPlan(steps=pending)

    # -- execution + self-heal ---------------------------------------------
    def _finalize(self, plan):
        for i, step in enumerate(plan.steps):
            self.store.event(self.task, "progress", f"正在执行第 {i + 1}/{len(plan.steps)} 步：{step.kind}")
            output = run_operation(self.store, self.task, step.model_dump())
            self.store.event(self.task, "output", output)
            emit_step(self.store, self.task, step.kind, output)

    def run(self, prompt):
        messages = [{"role": "system", "content": SYSTEM_PROMPT}, {"role": "user", "content": prompt}]
        plan = self._plan(messages)
        try:
            self._finalize(plan)
        except Exception as exc:
            if self._healed:
                raise
            self._healed = True
            messages.append({"role": "user", "content": f"执行中某步失败：{exc}。请结合已完成的输出修正操作，只给出尚未完成的剩余步骤。"})
            self._finalize(self._plan(messages))
        self.store.event(self.task, "message", self._final or "处理完成")
