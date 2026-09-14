"""Local Excel exchange for the explicitly supported workbook feature set."""
import colorsys
import csv
import posixpath
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles.colors import COLOR_INDEX as COLOR_INDEXED
from openpyxl.styles.numbers import is_date_format, is_datetime
from openpyxl.utils import range_boundaries, get_column_letter
from openpyxl.utils.datetime import to_excel, from_excel, WINDOWS_EPOCH

from .workbook_model import (MAX_CELLS, blank_workbook, blank_sheet, coordinate, position,
                             validate_snapshot, formula_issue, formula_errors)

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
BORDER_STYLES = {'thin': 1, 'medium': 2, 'dashed': 3, 'dotted': 4, 'thick': 5,
                 'double': 6, 'hair': 7, 'mediumDashed': 8, 'dashDot': 9,
                 'mediumDashDot': 10, 'dashDotDot': 11, 'mediumDashDotDot': 12, 'slantDashDot': 13}


def preflight(path):
    """Count XML cells before openpyxl allocates a large workbook."""
    count, issues = 0, []
    with ZipFile(path) as archive:
        names = archive.namelist()
        for prefix, message in [('xl/charts/', '图表'), ('xl/pivotTables/', '透视表'),
                                ('xl/externalLinks/', '外部链接'), ('xl/drawings/', '图片或绘图'),
                                ('xl/embeddings/', '嵌入对象'), ('xl/activeX/', 'ActiveX'),
                                ('xl/comments', '批注'), ('xl/printerSettings/', '打印机配置')]:
            if any(n.startswith(prefix) for n in names):
                issues.append(f'包含{message}，完整编辑暂不支持')
        if 'xl/connections.xml' in names or any('vbaProject' in n for n in names):
            issues.append('包含宏或外部数据连接')
        if 'xl/sharedStrings.xml' in names:
            with archive.open('xl/sharedStrings.xml') as stream:
                for _, elem in ET.iterparse(stream, events=('end',)):
                    if elem.tag == NS+'r':
                        issues.append('包含单元格内分段文字格式，完整编辑暂不支持')
                        break
                    elem.clear()
        for name in names:
            if not name.startswith('xl/worksheets/sheet') or not name.endswith('.xml'):
                continue
            with archive.open(name) as stream:
                for _, elem in ET.iterparse(stream, events=('end',)):
                    if elem.tag == NS + 'c':
                        if elem.find('.//'+NS+'r') is not None:
                            issues.append('包含单元格内分段文字格式，完整编辑暂不支持')
                        if elem.get('s') not in (None, '0') or len(elem):
                            count += 1
                        if count > MAX_CELLS:
                            return count, ['超过 200000 个有效单元格，请使用分页预览与全量处理']
                        if elem.find(NS + 'f') is not None and elem.find(NS + 'f').get('t') in {'array', 'dataTable'}:
                            issues.append('包含数组公式或数据表公式')
                        elem.clear()
                    elif elem.tag == NS + 'row':
                        elem.clear()
    return count, list(dict.fromkeys(issues))


def theme_colors(book):
    if not book.loaded_theme:
        return []
    root = ET.fromstring(book.loaded_theme)
    scheme = root.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}clrScheme')
    result = []
    if scheme is not None:
        for color in scheme:
            child = next(iter(color), None)
            result.append(child.get('lastClr', child.get('val', '000000')) if child is not None else '000000')
    return result


def rgb(color, theme, default='#000000'):
    if color is None:
        return default
    if color.type == 'rgb':
        base = color.rgb[-6:]
    elif color.type == 'indexed' and color.indexed < len(COLOR_INDEXED):
        base = COLOR_INDEXED[color.indexed][-6:]
    elif color.type == 'theme' and color.theme < len(theme):
        base = theme[color.theme][-6:]
    else:
        return default
    if color.tint:
        r, g, b = (int(base[i:i + 2], 16) / 255 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        l = l * (1 + color.tint) if color.tint < 0 else l * (1 - color.tint) + color.tint
        base = ''.join(f'{round(c * 255):02X}' for c in colorsys.hls_to_rgb(h, l, s))
    return '#' + base.upper()


def read_style(cell, theme):
    if not cell.has_style:
        return {}
    font, fill, align, border = cell.font, cell.fill, cell.alignment, cell.border
    style = {'font_name': font.name or 'Calibri', 'font_size': float(font.sz or 11),
             'bold': bool(font.b), 'italic': bool(font.i), 'underline': bool(font.u),
             'strike': bool(font.strike), 'font_color': rgb(font.color, theme),
             'num_format': cell.number_format or 'General'}
    if getattr(fill, 'patternType', None) == 'solid':
        style['bg_color'] = rgb(fill.fgColor, theme, '#FFFFFF')
    if align.horizontal in {'left', 'right', 'center', 'justify'}:
        style['align'] = align.horizontal
    if align.vertical:
        style['valign'] = 'center' if align.vertical == 'center' else align.vertical if align.vertical in {'top', 'bottom', 'justify'} else 'bottom'
    if align.wrap_text:
        style['text_wrap'] = True
    if align.text_rotation and align.text_rotation != 255:
        style['rotation'] = align.text_rotation if align.text_rotation <= 90 else 90 - align.text_rotation
    borders = {}
    for direction in ('left', 'right', 'top', 'bottom'):
        edge = getattr(border, direction)
        if edge and edge.style:
            borders[direction] = {'style': edge.style, 'color': rgb(edge.color, theme)}
    if borders:
        style['borders'] = borders
    return style


def typed_value(value, data_type=None):
    if isinstance(value, (datetime, date, time)):
        return {'value': to_excel(value, WINDOWS_EPOCH), 'type': 'date'}
    if data_type == 'e':
        return {'value': value, 'type': 'error'}
    if isinstance(value, bool):
        return {'value': value, 'type': 'boolean'}
    if isinstance(value, (float, int)):
        return {'value': value, 'type': 'number'}
    return {'value': str(value) if value is not None else None, 'type': 'string'}


def import_workbook(path, name=None):
    path = Path(path)
    result = blank_workbook(name or path.stem)
    issues = []
    if path.suffix == '.csv':
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                with path.open(encoding=encoding, newline='') as stream:
                    sheet = blank_sheet('数据')
                    for r, row in enumerate(csv.reader(stream)):
                        for c, value in enumerate(row):
                            if value:
                                sheet['cells'][coordinate(r, c)] = {'value': value, 'type': 'string'}
                        if len(sheet['cells']) > MAX_CELLS:
                            raise ValueError('完整编辑最多支持 200000 个有效单元格，请使用分页预览与全量处理')
                        sheet['row_count'] = max(sheet['row_count'], r + 1)
                        sheet['column_count'] = max(sheet['column_count'], len(row))
                result['sheets'] = [sheet]
                break
            except UnicodeDecodeError:
                if encoding == 'gb18030':
                    raise
        validate_snapshot(result)
        return result, issues
    if path.suffix == '.xls':
        import pandas as pd
        result['sheets'] = []
        with pd.ExcelFile(path, engine='calamine') as source:
            for name in source.sheet_names:
                frame = pd.read_excel(source, sheet_name=name, header=None, dtype=object, keep_default_na=False)
                sheet = blank_sheet(name)
                sheet.update(row_count=max(100, len(frame)), column_count=max(26, len(frame.columns)))
                for r, row in enumerate(frame.itertuples(index=False, name=None)):
                    for c, value in enumerate(row):
                        if value != '':
                            sheet['cells'][coordinate(r, c)] = typed_value(value)
                result['sheets'].append(sheet)
        validate_snapshot(result)
        return result, ['旧版 XLS 仅支持纯数据转换；原公式与格式不属于此转换范围']
    count, issues = preflight(path)
    if count > MAX_CELLS:
        raise ValueError(issues[0])
    book = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    # Retain numeric date serials, including Excel's fictitious 1900-02-29 (serial 60).
    serials = []
    with ZipFile(path) as archive:
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {r.get('Id'): r.get('Target') for r in relationships if r.get('Type', '').endswith('/worksheet')}
        workbook_xml = ET.fromstring(archive.read('xl/workbook.xml'))
        members = []
        for sheet_xml in workbook_xml.find(NS+'sheets'):
            target = targets.get(sheet_xml.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'))
            if target:
                members.append(posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join('xl', target)))
        for member in members:
            values = {}
            if member in archive.namelist():
                with archive.open(member) as stream:
                    for _, node in ET.iterparse(stream, events=('end',)):
                        if node.tag == NS+'c':
                            value = node.find(NS+'v')
                            if node.get('t', 'n') == 'n' and value is not None and value.text:
                                values[node.get('r')] = value.text
                            node.clear()
                        elif node.tag == NS+'row': node.clear()
            serials.append(values)
    try:
        theme = theme_colors(book)
        result['date1904'] = book.epoch.year == 1904
        result['sheets'] = []
        if book.defined_names:
            issues.append('包含命名区域，首版编辑器暂不支持')
        for sheet_index, source in enumerate(book.worksheets):
            sheet = blank_sheet(source.title)
            sheet.update(row_count=max(100, source.max_row), column_count=max(26, source.max_column), hidden=source.sheet_state != 'visible')
            if source.tables or source.conditional_formatting or source.data_validations.count:
                issues.append(f'{source.title}：包含结构化表、条件格式或数据验证')
            if source.protection.sheet:
                issues.append(f'{source.title}：包含工作表保护')
            if source.print_area or source.print_title_rows or source.print_title_cols or source.row_breaks.count or source.col_breaks.count:
                issues.append(f'{source.title}：包含打印区域、重复标题或分页配置')
            if source.page_setup.orientation or source.page_setup.paperSize or source.page_setup.scale or source.page_setup.fitToWidth or source.page_setup.fitToHeight:
                issues.append(f'{source.title}：包含自定义打印布局')
            if any(d.outlineLevel or d.collapsed for d in [*source.row_dimensions.values(), *source.column_dimensions.values()]):
                issues.append(f'{source.title}：包含行列分组或折叠结构')
            for (r, c), original in source._cells.items():
                if original.value is None and not original.has_style:
                    continue
                cell = typed_value(original.value, original.data_type)
                if original.data_type == 'f':
                    if not isinstance(original.value, str):
                        issues.append(f'{source.title}!{original.coordinate}：数组或数据表公式不支持')
                        cell = typed_value(cached[source.title].cell(r, c).value)
                    else:
                        issue = formula_issue(original.value)
                        if issue:
                            issues.append(f'{source.title}!{original.coordinate}：{issue}')
                        value_cell = cached[source.title].cell(r, c)
                        cell = typed_value(value_cell.value, value_cell.data_type)
                        cell['formula'] = original.value
                        cell['result_state'] = 'pending' if value_cell.value is None else 'error' if value_cell.data_type == 'e' else 'ready'
                if original.hyperlink or original.comment:
                    issues.append(f'{source.title}：包含超链接或批注')
                if original.data_type == 'e':
                    issues.append(f'{source.title}：包含直接存储的 Excel 错误值，请先在 Excel 中处理或提取纯数据副本')
                if original.font.vertAlign or original.font.u not in {None,'single'} or original.border.diagonalUp or original.border.diagonalDown or original.alignment.shrink_to_fit or original.alignment.indent:
                    issues.append(f'{source.title}：包含上下标、特殊下划线、对角边框或特殊文字排版')
                if getattr(original.fill, 'patternType', 'gradient') not in {None, 'solid'} or original.alignment.text_rotation == 255:
                    issues.append(f'{source.title}：包含特殊填充或竖排文字')
                style = read_style(original, theme)
                if is_date_format(original.number_format) and original.coordinate in serials[sheet_index]:
                    serial = float(serials[sheet_index][original.coordinate])
                    if result['date1904'] and is_datetime(original.number_format) != 'time':
                        serial += 1462
                    cell.update(value=serial, type='date')
                if style:
                    cell['style'] = style
                sheet['cells'][coordinate(r - 1, c - 1)] = cell
            for r, dim in source.row_dimensions.items():
                props = {'hidden': bool(dim.hidden)}
                if dim.height is not None:
                    props['height'] = dim.height * 4 / 3
                if dim.has_style:
                    props['style'] = read_style(dim, theme)
                sheet['rows'][str(r - 1)] = props
            for _, dim in source.column_dimensions.items():
                for c in range(dim.min, dim.max + 1):
                    props = {'width': float(dim.width or 13) * 7 + 5, 'hidden': bool(dim.hidden)}
                    if dim.has_style:
                        props['style'] = read_style(dim, theme)
                    sheet['columns'][str(c - 1)] = props
            sheet['merges'] = [dict(r0=m.min_row - 1, c0=m.min_col - 1, r1=m.max_row - 1, c1=m.max_col - 1) for m in source.merged_cells.ranges]
            if source.freeze_panes:
                from openpyxl.utils.cell import coordinate_to_tuple
                r, c = coordinate_to_tuple(source.freeze_panes)
                sheet['freeze'] = {'row': r - 1, 'column': c - 1}
            if source.auto_filter.ref:
                c0, r0, c1, r1 = range_boundaries(source.auto_filter.ref)
                columns = []
                for f in source.auto_filter.filterColumn:
                    if f.filters and not f.filters.dateGroupItem:
                        columns.append({'column': c0 - 1 + f.colId, 'values': list(f.filters.filter), 'blank': bool(f.filters.blank)})
                    else:
                        issues.append(f'{source.title}：包含暂不支持的高级筛选条件')
                hidden_rows = []
                for r in range(r0 + 1, r1 + 1):
                    if not columns or not source.row_dimensions.get(r) or not source.row_dimensions[r].hidden:
                        continue
                    passed = True
                    for entry in columns:
                        value = cached[source.title].cell(r, entry['column'] + 1).value
                        text = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
                        passed &= (value is None and entry['blank']) or text in [str(v) for v in entry['values']]
                    if not passed:
                        hidden_rows.append(r - 1)
                        sheet['rows'][str(r - 1)]['hidden'] = False
                sheet['filter'] = {'range': dict(r0=r0 - 1, c0=c0 - 1, r1=r1 - 1, c1=c1 - 1), 'columns': columns, 'hidden_rows': hidden_rows}
            result['sheets'].append(sheet)
        validate_snapshot(result, allow_unsupported=True)
        return result, list(dict.fromkeys(issues))
    finally:
        book.close(); cached.close()


def writer_style(style):
    value = {k: v for k, v in style.items() if k not in {'borders', 'rotation'}}
    if 'strike' in value:
        value['font_strikeout'] = value.pop('strike')
    if value.get('valign') == 'center':
        value['valign'] = 'vcenter'
    if value.get('align') == 'general':
        value.pop('align')
    if 'rotation' in style:
        value['rotation'] = style['rotation']
    for direction, edge in style.get('borders', {}).items():
        value[direction] = BORDER_STYLES[edge['style']]
        value[direction + '_color'] = edge.get('color', '#000000')
    return value


def export_workbook(snapshot, path, *, require_calculated=True):
    import json
    import xlsxwriter
    validate_snapshot(snapshot)
    if require_calculated and formula_errors(snapshot):
        raise ValueError('请完成公式重算并修复错误后导出：' + '；'.join(formula_errors(snapshot)[:5]))
    with xlsxwriter.Workbook(path, {'strings_to_formulas': False, 'strings_to_urls': False, 'date_1904': snapshot['date1904']}) as book:
        formats = {}
        def fmt(style):
            if not style:
                return None
            key = json.dumps(style, sort_keys=True)
            if key not in formats:
                formats[key] = book.add_format(writer_style(style))
            return formats[key]
        for data in snapshot['sheets']:
            sheet = book.add_worksheet(data['name'])
            if data.get('hidden'):
                sheet.hide()
            for r, props in data.get('rows', {}).items():
                sheet.set_row(int(r), props['height'] * .75 if 'height' in props else None, fmt(props.get('style')), {'hidden': props.get('hidden', False)})
            for c, props in data.get('columns', {}).items():
                sheet.set_column_pixels(int(c), int(c), props.get('width', 96), fmt(props.get('style')), {'hidden': props.get('hidden', False)})
            for m in data.get('merges', []):
                sheet.merge_range(m['r0'], m['c0'], m['r1'], m['c1'], '')
            for key, cell in data['cells'].items():
                r, c = position(key)
                value, style = cell.get('value'), fmt(cell.get('style'))
                inherited = {**data.get('columns', {}).get(str(c), {}).get('style', {}), **data.get('rows', {}).get(str(r), {}).get('style', {}), **cell.get('style', {})}
                number_format = inherited.get('num_format', 'yyyy-mm-dd' if cell.get('type') == 'date' else '')
                if cell.get('formula'):
                    if cell.get('result_state') != 'ready':
                        raise ValueError(f"{data['name']}!{get_column_letter(c + 1)}{r + 1} 公式尚未计算")
                    cache = value
                    if snapshot['date1904'] and is_date_format(number_format) and is_datetime(number_format) != 'time' and isinstance(cache, (int, float)):
                        cache -= 1462
                    sheet.write_formula(r, c, cell['formula'], style, cache)
                elif value is None:
                    sheet.write_blank(r, c, None, style)
                elif cell.get('type') == 'date':
                    serial = value - 1462 if snapshot['date1904'] and is_datetime(number_format) != 'time' else value
                    sheet.write_number(r, c, serial, style or fmt({'num_format': 'yyyy-mm-dd'}))
                elif isinstance(value, bool):
                    sheet.write_boolean(r, c, value, style)
                elif isinstance(value, (int, float)):
                    # Univer uses the 1900 serial system, including manually edited dates.
                    is_date = is_date_format(number_format)
                    sheet.write_number(r, c, value - 1462 if snapshot['date1904'] and is_date and is_datetime(number_format) != 'time' else value, style)
                else:
                    sheet.write_string(r, c, str(value), style)
            freeze = data.get('freeze', {})
            sheet.freeze_panes(freeze.get('row', 0), freeze.get('column', 0))
            filters = data.get('filter')
            if filters:
                area = filters['range']
                sheet.autofilter(area['r0'], area['c0'], area['r1'], area['c1'])
                for entry in filters.get('columns', []):
                    values = list(entry.get('values', [])) + (['Blanks'] if entry.get('blank') else [])
                    if values:
                        sheet.filter_column_list(entry['column'], values)
                for row in filters.get('hidden_rows', []):
                    props = data.get('rows', {}).get(str(row), {})
                    sheet.set_row(row, props['height']*.75 if 'height' in props else None, fmt(props.get('style')), {'hidden': True})
    check = load_workbook(path, read_only=True)
    check.close()
